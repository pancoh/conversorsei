"""Bloco de assinatura: as linhas centralizadas do fim do documento.

Saem como Tabela_Texto_Centralizado, sem a margem de 6 pt de Texto_Centralizado, que
espaçava demais o "[Assinado eletronicamente]", o nome e o cargo.
"""
from __future__ import annotations

import io
import re
import zipfile

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from conversorsei.web import converter_documento_memoria

RE_PARAGRAFO = re.compile(r'<p class="([^"]+)">(.*?)</p>', re.DOTALL)
CENTRO = WD_ALIGN_PARAGRAPH.CENTER


def paragrafos(conteudo: bytes, nome: str = "t.docx") -> list[tuple[str, str]]:
    """(classe, HTML interno) de cada parágrafo, na ordem."""
    res = converter_documento_memoria(nome, conteudo, so_corpo=True)
    assert res["sucesso"], res["erros"]
    return RE_PARAGRAFO.findall("\n".join(a["conteudo"] for a in res["arquivos"]))


def _p(doc, texto: str, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, **fonte):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    if texto:
        run = p.add_run(texto)
        for atributo, valor in fonte.items():
            setattr(run.font, atributo, valor)
    return p


def _salvar(doc) -> bytes:
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def nota_com_assinatura() -> bytes:
    doc = docx.Document()
    _p(doc, "Ministério das Cidades", CENTRO, bold=True)
    _p(doc, "Trata-se de análise da proposição, com os fundamentos a seguir.")
    _p(doc, "Brasília, na data da assinatura.", WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "")
    _p(doc, "[Assinado eletronicamente]", CENTRO)
    _p(doc, "FULANO DE TAL", CENTRO, bold=True)
    _p(doc, "Assessor Técnico", CENTRO)
    _p(doc, "")
    # Variação da marca, que sai padronizada
    _p(doc, "(assinado eletronicamente)", CENTRO)
    # Caixa alta pelo atributo da fonte, e não no texto
    _p(doc, "Ciclano & Beltrano", CENTRO, bold=True, all_caps=True)
    _p(doc, "Coordenador-Geral", CENTRO)
    return _salvar(doc)


def test_linhas_centralizadas_do_fim_viram_bloco_de_assinatura():
    obtido = paragrafos(nota_com_assinatura())
    classes = [c for c, _ in obtido]

    assert classes[:3] == ["Texto_Centralizado", "Texto_Justificado", "Texto_Alinhado_Direita"]
    assert set(classes[3:]) == {"Tabela_Texto_Centralizado"}


def test_marca_em_italico_e_linha_em_branco_entre_as_assinaturas():
    """A marca sai padronizada e em itálico; antes de cada marca, menos a primeira, uma
    linha em branco separa as assinaturas."""
    internos = [i for _, i in paragrafos(nota_com_assinatura())[3:]]
    assert internos == [
        "<em>[Assinado eletronicamente]</em>",
        "<strong>FULANO DE TAL</strong>",
        "Assessor Técnico",
        "&nbsp;",
        "<em>[Assinado eletronicamente]</em>",
        "<strong>CICLANO &amp; BELTRANO</strong>",
        "Coordenador-Geral",
    ]


def test_uma_assinatura_so_nao_tem_linha_em_branco():
    doc = docx.Document()
    _p(doc, "Texto do corpo.")
    for texto in ("Assinado eletronicamente", "Fulano de Tal", "Assessor Técnico"):
        _p(doc, texto, CENTRO)
    internos = [i for _, i in paragrafos(_salvar(doc))]
    assert internos[1:] == ["<em>[Assinado eletronicamente]</em>", "Fulano de Tal", "Assessor Técnico"]


def test_assinaturas_sem_marca_nao_sao_separadas():
    """Sem a marca, não dá para saber onde uma assinatura termina e a outra começa."""
    doc = docx.Document()
    _p(doc, "Texto do corpo.")
    for texto in ("Fulano de Tal", "Assessor Técnico", "Ciclano de Tal", "Coordenador"):
        _p(doc, texto, CENTRO)
    assert "&nbsp;" not in [i for _, i in paragrafos(_salvar(doc))]


def test_marca_no_meio_do_texto_nao_muda():
    doc = docx.Document()
    _p(doc, "Assinado eletronicamente")
    _p(doc, "Texto do corpo.")
    assert paragrafos(_salvar(doc))[0] == ("Texto_Justificado", "Assinado eletronicamente")


def test_nome_da_assinatura_mantem_negrito_e_caixa_alta_no_texto():
    """A classe da assinatura não aplica negrito nem caixa alta: o texto passa a levá-los."""
    internos = [i for _, i in paragrafos(nota_com_assinatura())]
    assert "<strong>FULANO DE TAL</strong>" in internos
    # A entidade do & fica intacta ao passar para caixa alta
    assert "<strong>CICLANO &amp; BELTRANO</strong>" in internos


def test_paragrafo_centralizado_longo_no_fim_nao_e_assinatura():
    doc = docx.Document()
    _p(doc, "Texto do corpo.")
    _p(doc, "Fulano de Tal", CENTRO)
    _p(doc, "Este parágrafo centralizado é longo demais para ser linha de assinatura de alguém.", CENTRO)
    classes = [c for c, _ in paragrafos(_salvar(doc))]
    assert classes == ["Texto_Justificado", "Texto_Centralizado", "Texto_Centralizado"]


def test_uma_linha_centralizada_sozinha_no_fim_nao_e_assinatura():
    """Legenda ou fecho: assinatura tem ao menos nome e cargo."""
    doc = docx.Document()
    _p(doc, "Texto do corpo.")
    _p(doc, "Figura 1 - Mapa da rede", CENTRO)
    assert paragrafos(_salvar(doc))[-1][0] == "Texto_Centralizado"


def test_documento_todo_centralizado_nao_tem_bloco():
    doc = docx.Document()
    for texto in ("Certificado", "Fulano de Tal", "Participou do curso"):
        _p(doc, texto, CENTRO)
    assert {c for c, _ in paragrafos(_salvar(doc))} == {"Texto_Centralizado"}


def test_tabela_no_fim_encerra_a_busca():
    doc = docx.Document()
    _p(doc, "Texto do corpo.")
    _p(doc, "Fulano de Tal", CENTRO)
    doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Valor")
    obtido = dict((i, c) for c, i in paragrafos(_salvar(doc)))
    assert obtido["Fulano de Tal"] == "Texto_Centralizado"


def test_estilo_com_nome_de_classe_no_fim_continua_valendo():
    """Estilo com o nome de uma classe é pedido explícito, também no fim do documento."""
    markdown = (
        "Texto do corpo.\n\n"
        '<p class="Texto_Alinhado_Direita">Brasília, 1 de janeiro de 2026.</p>\n\n'
        '<p class="Texto_Centralizado">Fulano de Tal</p>\n'
    )
    classes = [c for c, _ in paragrafos(markdown.encode(), "t.md")]
    assert classes[-1] == "Texto_Centralizado"


CONTENT_ODT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:automatic-styles>
  <style:style style:name="P1" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center"/>
  </style:style>
  <style:style style:name="P2" style:family="paragraph">
   <style:paragraph-properties fo:text-align="end"/>
  </style:style>
  <style:style style:name="P3" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center"/>
   <style:text-properties fo:font-weight="bold" fo:text-transform="uppercase"/>
  </style:style>
  <style:style style:name="P4" style:family="paragraph">
   <style:paragraph-properties fo:text-align="start"/>
  </style:style>
 </office:automatic-styles>
 <office:body>
  <office:text>
   <text:p text:style-name="P1">Ministério das Cidades</text:p>
   <text:p text:style-name="P4">Texto do corpo, com uma frase comum.</text:p>
   <text:p text:style-name="P2">Brasília, na data da assinatura.</text:p>
   <text:p text:style-name="P1">[Assinado eletronicamente]</text:p>
   <text:p text:style-name="P3">Fulano de Tal</text:p>
   <text:p text:style-name="P1">Assessor Técnico</text:p>
  </office:text>
 </office:body>
</office:document-content>
"""


def _odt(content: str) -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
    return saida.getvalue()


def test_odt_leva_o_alinhamento_e_o_bloco_de_assinatura_como_o_word():
    """Centralizado e à direita valem no ODT; à esquerda (em geral herdado) segue justificado."""
    assert paragrafos(_odt(CONTENT_ODT), "t.odt") == [
        ("Texto_Centralizado", "Ministério das Cidades"),
        ("Texto_Justificado", "Texto do corpo, com uma frase comum."),
        ("Texto_Alinhado_Direita", "Brasília, na data da assinatura."),
        ("Tabela_Texto_Centralizado", "<em>[Assinado eletronicamente]</em>"),
        # Negrito e caixa alta do estilo passam para o texto
        ("Tabela_Texto_Centralizado", "<strong>FULANO DE TAL</strong>"),
        ("Tabela_Texto_Centralizado", "Assessor Técnico"),
    ]


def test_odt_com_uma_linha_centralizada_no_fim_mantem_texto_centralizado():
    content = CONTENT_ODT.replace(
        '<text:p text:style-name="P1">[Assinado eletronicamente]</text:p>\n', ""
    ).replace('<text:p text:style-name="P3">Fulano de Tal</text:p>\n', "")
    assert paragrafos(_odt(content), "t.odt")[-1] == ("Texto_Centralizado", "Assessor Técnico")


def test_odt_com_duas_assinaturas_tem_linha_em_branco_entre_elas():
    segunda = (
        '<text:p text:style-name="P1">[Assinado eletronicamente]</text:p>\n'
        '   <text:p text:style-name="P3">Ciclano de Tal</text:p>\n'
        '   <text:p text:style-name="P1">Coordenador-Geral</text:p>\n'
        "  </office:text>"
    )
    content = CONTENT_ODT.replace("  </office:text>", "   " + segunda)
    internos = [i for _, i in paragrafos(_odt(content), "t.odt")[3:]]
    assert internos == [
        "<em>[Assinado eletronicamente]</em>",
        "<strong>FULANO DE TAL</strong>",
        "Assessor Técnico",
        "&nbsp;",
        "<em>[Assinado eletronicamente]</em>",
        "<strong>CICLANO DE TAL</strong>",
        "Coordenador-Geral",
    ]


# Sem centralização, a marca de assinatura no fim do documento abre o bloco: o texto
# colado de um chat ou de um .txt não tem alinhamento, mas a marca, o nome e o cargo
# dizem o que é a assinatura
TEXTO_SEM_ALINHAMENTO = """Sobre o assunto, informa-se que esta unidade não possui área com essa finalidade.

Atenciosamente,

[assinado eletronicamente]
FULANO DE TAL
Coordenador-Geral
"""


def test_marca_no_fim_abre_o_bloco_mesmo_sem_centralizacao():
    for nome in ("t.txt", "t.md"):
        classes = [c for c, _ in paragrafos(TEXTO_SEM_ALINHAMENTO.encode(), nome)]
        assert classes == ["Texto_Justificado", "Texto_Justificado"] + ["Tabela_Texto_Centralizado"] * 3, nome

    html = "".join(f"<p>{linha}</p>" for linha in TEXTO_SEM_ALINHAMENTO.splitlines() if linha)
    corpo = paragrafos(html.encode(), "texto_colado.html")
    assert [c for c, _ in corpo][-3:] == ["Tabela_Texto_Centralizado"] * 3
    assert corpo[-4] == ("Texto_Justificado", "Atenciosamente,")
    assert corpo[-3][1] == "<em>[Assinado eletronicamente]</em>"


def test_duas_assinaturas_sem_centralizacao_saem_separadas():
    texto = TEXTO_SEM_ALINHAMENTO + "[assinado eletronicamente]\nCICLANO DE TAL\nDiretor\n"
    corpo = paragrafos(texto.encode(), "t.txt")
    assert [c for c, _ in corpo[-7:]] == ["Tabela_Texto_Centralizado"] * 7
    assert corpo[-4][1] == "&nbsp;"


def test_sem_marca_ou_com_texto_depois_dela_nada_muda():
    sem_marca = TEXTO_SEM_ALINHAMENTO.replace("[assinado eletronicamente]\n", "")
    assert "Tabela_Texto_Centralizado" not in {c for c, _ in paragrafos(sem_marca.encode(), "t.txt")}

    # A marca seguida direto de um parágrafo comprido: falta nome e cargo para ser assinatura
    so_marca = TEXTO_SEM_ALINHAMENTO.replace("FULANO DE TAL\nCoordenador-Geral\n", "") + (
        "Este parágrafo final é comprido demais para ser nome ou cargo, e por isso a marca acima não abre bloco.\n"
    )
    assert "Tabela_Texto_Centralizado" not in {c for c, _ in paragrafos(so_marca.encode(), "t.txt")}


def test_assinatura_no_meio_do_documento_tambem_e_bloco():
    """A nota assinada pelo coordenador e, depois do "De acordo.", pela diretora.

    Só a assinatura do fim era reconhecida; a do meio saía como texto centralizado comum.
    """
    doc = docx.Document()
    _p(doc, "Encaminha-se a Nota Técnica para análise.")
    _p(doc, "Respeitosamente,")
    _p(doc, "[assinado eletronicamente]", CENTRO)
    _p(doc, "FULANO DE TAL", CENTRO)
    _p(doc, "Coordenador-Geral de Gestão e Monitoramento do Programa", CENTRO)
    _p(doc, "CGPAC/DMP/SE/MCID", CENTRO)
    _p(doc, "De acordo. À Assessoria Especial de Controle Interno para interlocução com a CGU.")
    _p(doc, "[assinado eletronicamente]", CENTRO)
    _p(doc, "CICLANA DE TAL", CENTRO)
    _p(doc, "Diretora de Monitoramento de Programas", CENTRO)
    classes = [c for c, _ in paragrafos(_salvar(doc))]

    assert classes == (
        ["Texto_Justificado"] * 2 + ["Tabela_Texto_Centralizado"] * 4
        + ["Texto_Justificado"] + ["Tabela_Texto_Centralizado"] * 3
    )
    marcas = [t for _, t in paragrafos(_salvar(doc)) if "Assinado" in t]
    assert marcas == ["<em>[Assinado eletronicamente]</em>"] * 2


def test_assinatura_no_meio_sem_centralizacao_para_antes_do_de_acordo():
    """Sem centralização, o bloco do meio para na linha que termina em pontuação."""
    texto = TEXTO_SEM_ALINHAMENTO + "De acordo.\n\n[assinado eletronicamente]\nCICLANA DE TAL\nDiretora\n"
    corpo = paragrafos(texto.encode(), "t.txt")
    assert [c for c, _ in corpo] == (
        ["Texto_Justificado"] * 2 + ["Tabela_Texto_Centralizado"] * 3
        + ["Texto_Justificado"] + ["Tabela_Texto_Centralizado"] * 3
    )
    assert corpo[5][1] == "De acordo."


def test_linhas_centralizadas_no_meio_sem_marca_nao_sao_assinatura():
    doc = docx.Document()
    _p(doc, "Texto de abertura do documento.")
    _p(doc, "QUADRO RESUMO", CENTRO)
    _p(doc, "Situação em 2026", CENTRO)
    _p(doc, "Texto que continua o documento depois do destaque centralizado.")
    assert "Tabela_Texto_Centralizado" not in {c for c, _ in paragrafos(_salvar(doc))}


def test_linhas_de_assinatura_pela_marca():
    from conversorsei.formatacao import linhas_de_assinatura

    curtas = ["Texto.", "Atenciosamente,", "[assinado eletronicamente]", "NOME", "Cargo"]
    textos = [None, *curtas[1:]]
    assert linhas_de_assinatura([False] * 5, textos) == 3
    # Só a marca, sem nome: menos linhas que o mínimo
    assert linhas_de_assinatura([False] * 3, [None, "Atenciosamente,", "[assinado eletronicamente]"]) == 0
    # A regra da centralização continua valendo, e vale a maior das duas
    assert linhas_de_assinatura([False, True, True, True, True], textos) == 4


def test_assinatura_com_quebras_de_linha_num_paragrafo_so():
    """O chat escreve a marca, o nome e o cargo num parágrafo, separados por <br>."""
    fim = "<p>Diante do exposto, submeto o processo.</p>"
    for assinatura in (
        "<p>Atenciosamente,</p><p>[assinado eletronicamente]<br>FULANO DE TAL<br>Coordenador-Geral</p>",
        "<p>Atenciosamente,<br><br>[assinado eletronicamente]<br>FULANO DE TAL<br>Coordenador-Geral</p>",
    ):
        corpo = paragrafos((fim + assinatura).encode(), "texto_colado.html")
        assert corpo == [
            ("Texto_Justificado", "Diante do exposto, submeto o processo."),
            ("Texto_Justificado", "Atenciosamente,"),
            ("Tabela_Texto_Centralizado", "<em>[Assinado eletronicamente]</em>"),
            ("Tabela_Texto_Centralizado", "FULANO DE TAL"),
            ("Tabela_Texto_Centralizado", "Coordenador-Geral"),
        ]


def test_quebra_de_linha_sem_marca_continua_no_mesmo_paragrafo():
    corpo = paragrafos(b"<p>Texto com linha<br>quebrada no meio.</p>", "texto_colado.html")
    assert corpo == [("Texto_Justificado", "Texto com linha<br />quebrada no meio.")]
