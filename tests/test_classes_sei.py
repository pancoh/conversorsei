"""Cobertura das 34 classes do CSS institucional do SEI.

O documento de referência exercita cada classe pelo caminho que o usuário usaria no Word:
a formatação, quando formatacao.py tem regra para ela (caixa alta, riscado, fundo cinza,
fonte da tabela), e o nome do estilo nas demais. Se uma regra mudar e deixar uma classe
sem caminho, o teste aponta qual.
"""
from __future__ import annotations

import html
import io
import re
import zipfile
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

from conversorsei.formatacao import classes_sei
from conversorsei.odt_converter import extrair_markdown_odt
from conversorsei.web import converter_documento_memoria

RE_PARAGRAFO = re.compile(r'<(p|li) class="([^"]+)">(.*?)</\1>', re.DOTALL)

# Classes que só o nome do estilo alcança. Ementa, bloco de código, endereçamento e caixa
# alta justificada tinham regra por medida (recuo, fonte, espaçamento), retirada porque
# errava em documentos comuns (assinatura recuada, documento inteiro em Courier)
SO_PELO_NOME = (
    "Texto_Citação",
    "Texto_Justificado_Recuo_Primeira_Linha2",
    "Fonte_Calibri",
    "Texto_Ementa",
    "Texto_Mono_Espaçado",
    "Texto_Justificado_Maiusculas",
    "Texto_Alinhado_Esquerda_Espaçamento_Simples",
    "Texto_Alinhado_Esquerda_Espacamento_Simples_Maiusc",
)


def classes_por_texto(html_corpo: str) -> dict[str, str]:
    """Mapeia o texto visível de cada parágrafo à classe que ele recebeu."""
    return {
        html.unescape(re.sub(r"<[^>]+>", "", interno)).strip(): classe
        for _, classe, interno in RE_PARAGRAFO.findall(html_corpo)
    }


def _definir_lista(doc, num_id: int, formato: str) -> None:
    """Acrescenta ao numbering.xml uma lista de um só nível no formato pedido."""
    numbering = doc.part.numbering_part.element
    numbering.insert(
        0,
        parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num_id}">'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{formato}"/></w:lvl>'
            "</w:abstractNum>"
        ),
    )
    numbering.append(
        parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>')
    )


def _numerar(p, num_id: int) -> None:
    numPr = OxmlElement("w:numPr")
    for tag, valor in (("w:ilvl", "0"), ("w:numId", str(num_id))):
        el = OxmlElement(tag)
        el.set(qn("w:val"), valor)
        numPr.append(el)
    p._p.get_or_add_pPr().append(numPr)


def _paragrafo(doc, texto: str, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, **fonte):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    run = p.add_run(texto)
    for atributo, valor in fonte.items():
        setattr(run.font, atributo, valor)
    return p


def _estilo(doc, nome: str) -> None:
    doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["Normal"]


def _celula(celula, texto: str, alinhamento, tamanho: float | None = None) -> None:
    p = celula.paragraphs[0]
    p.alignment = alinhamento
    run = p.add_run(texto)
    if tamanho:
        run.font.size = Pt(tamanho)


# (texto, classe esperada) de cada parágrafo do documento de referência, na ordem
ESPERADO_DOCX = [
    ("PORTARIA Nº 123, DE 25 DE SETEMBRO DE 2026", "Texto_Centralizado_Maiusculas_Negrito"),
    ("MINISTÉRIO DA INTEGRAÇÃO", "Texto_Centralizado_Maiusculas"),
    ("Dispõe sobre a conversão de documentos para o editor do SEI.", "Texto_Ementa"),
    ("O SECRETÁRIO, no uso de suas atribuições, resolve:", "Texto_Justificado"),
    ("R E S O L V E:", "Texto_Espaco_Duplo_Recuo_Primeira_Linha"),
    ("Contexto", "Item_Nivel1"),
    ("Antecedentes", "Item_Nivel2"),
    ("Histórico", "Item_Nivel3"),
    ("Registros", "Item_Nivel4"),
    ("Trata-se de pedido de análise, com base no processo, apresentado pela unidade.", "Paragrafo_Numerado_Nivel1"),
    ("Em complemento, observa-se que o pedido veio instruído, com os documentos exigidos.",
     "Paragrafo_Numerado_Nivel2"),
    ("Registre-se, ainda, que a instrução foi conferida pela área técnica, sem ressalvas.",
     "Paragrafo_Numerado_Nivel3"),
    ("inciso em algarismo romano;", "Item_Inciso_Romano"),
    ("alínea em letra;", "Item_Alinea_Letra"),
    ("item com marcador", "Texto_Justificado"),
    ("Art. 1º Texto citado de norma, pelo estilo Citação do Word.", "Citação"),
    ("Texto de citação pelo estilo institucional.", "Texto_Citação"),
    ("Art. 2º Dispositivo revogado.", "Tachado"),
    ("resultado = converter(documento)", "Texto_Mono_Espaçado"),
    ("Destaque com fundo cinza", "Texto_Fundo_Cinza_Negrito"),
    ("DESTAQUE EM CAIXA ALTA", "Texto_Fundo_Cinza_Maiusculas_Negrito"),
    ("TEXTO JUSTIFICADO EM CAIXA ALTA SEM NEGRITO", "Texto_Justificado_Maiusculas"),
    ("Parágrafo com recuo na primeira linha.", "Texto_Justificado_Recuo_Primeira_Linha"),
    ("Parágrafo com recuo pelo estilo alternativo.", "Texto_Justificado_Recuo_Primeira_Linha2"),
    ("Texto centralizado comum", "Texto_Centralizado"),
    ("Brasília, 25 de setembro de 2026.", "Texto_Alinhado_Direita"),
    ("Texto alinhado à esquerda com espaço entre parágrafos.", "Texto_Alinhado_Esquerda"),
    ("Ao Senhor", "Texto_Alinhado_Esquerda_Espaçamento_Simples"),
    ("FULANO DE TAL", "Texto_Alinhado_Esquerda_Espacamento_Simples_Maiusc"),
    ("Parágrafo só com a fonte institucional.", "Fonte_Calibri"),
    ("Coluna", "Tabela_Texto_10"),
    ("Valor", "Tabela_Texto_10"),
    ("nota de rodapé da tabela", "Tabela_Texto_8"),
    ("12,5", "Tabela_Texto_Alinhado_Direita"),
    ("Descrição", "Tabela_Texto_Alinhado_Esquerda"),
    ("Sim", "Tabela_Texto_Centralizado"),
]


def montar_docx_todas_as_classes() -> bytes:
    """Documento Word com um parágrafo para cada classe do SEI, pelo caminho que a produz."""
    doc = docx.Document()
    for nome in SO_PELO_NOME:
        _estilo(doc, nome)
    _definir_lista(doc, 90, "upperRoman")
    _definir_lista(doc, 91, "lowerLetter")
    centro, direita, esquerda = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT
    textos = iter(texto for texto, _ in ESPERADO_DOCX)

    _paragrafo(doc, next(textos), centro, bold=True)
    _paragrafo(doc, next(textos), centro)
    doc.add_paragraph(next(textos), style="Texto_Ementa")
    _paragrafo(doc, next(textos))
    p = _paragrafo(doc, next(textos), bold=True)
    espacamento = OxmlElement("w:spacing")
    espacamento.set(qn("w:val"), "40")
    p.runs[0]._r.get_or_add_rPr().append(espacamento)

    for nivel in range(1, 5):
        doc.add_heading(next(textos), level=nivel)
    for numero in ("1.", "1.1.", "1.1.1."):
        _paragrafo(doc, f"{numero} {next(textos)}")

    _numerar(_paragrafo(doc, next(textos)), 90)
    _numerar(_paragrafo(doc, next(textos)), 91)
    doc.add_paragraph(next(textos), style="List Bullet")

    doc.add_paragraph(next(textos), style="Quote")
    doc.add_paragraph(next(textos), style="Texto_Citação")
    _paragrafo(doc, next(textos), strike=True)
    doc.add_paragraph(next(textos), style="Texto_Mono_Espaçado")
    for _ in range(2):
        p = _paragrafo(doc, next(textos), bold=True)
        sombra = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="E6E6E6"/>')
        p._p.get_or_add_pPr().append(sombra)
    doc.add_paragraph(next(textos), style="Texto_Justificado_Maiusculas")
    _paragrafo(doc, next(textos)).paragraph_format.first_line_indent = Cm(2.5)
    doc.add_paragraph(next(textos), style="Texto_Justificado_Recuo_Primeira_Linha2")
    _paragrafo(doc, next(textos), centro)
    _paragrafo(doc, next(textos), direita)
    _paragrafo(doc, next(textos), esquerda)
    doc.add_paragraph(next(textos), style="Texto_Alinhado_Esquerda_Espaçamento_Simples")
    doc.add_paragraph(next(textos), style="Texto_Alinhado_Esquerda_Espacamento_Simples_Maiusc")
    doc.add_paragraph(next(textos), style="Fonte_Calibri")

    tabela = doc.add_table(rows=3, cols=2)
    _celula(tabela.cell(0, 0), next(textos), centro, 10)
    _celula(tabela.cell(0, 1), next(textos), centro, 10)
    _celula(tabela.cell(1, 0), next(textos), esquerda, 8)
    _celula(tabela.cell(1, 1), next(textos), direita)
    _celula(tabela.cell(2, 0), next(textos), esquerda)
    _celula(tabela.cell(2, 1), next(textos), centro)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def test_documento_de_referencia_usa_todas_as_classes_do_sei():
    resultado = converter_documento_memoria("todas_as_classes.docx", montar_docx_todas_as_classes())

    assert resultado["sucesso"], resultado["erros"]
    assert resultado["avisos"] == []
    assert len(resultado["arquivos"]) == 1
    corpo = resultado["arquivos"][0]["conteudo"]
    obtido = classes_por_texto(corpo)

    divergencias = {texto: (obtido.get(texto), classe) for texto, classe in ESPERADO_DOCX if obtido.get(texto) != classe}
    assert divergencias == {}, "texto: (obtida, esperada)"
    assert set(re.findall(r'class="([^"]+)"', corpo)) == set(classes_sei())


def test_markdown_tem_marcacao_para_as_classes_proprias():
    md = (
        "# NOTA TÉCNICA\n\n"
        "> Art. 1º Texto citado.\n\n"
        "```\ncodigo  *literal*\n  indentado\n```\n\n"
        "Texto com ~~trecho riscado~~ no meio.\n\n"
        "~~Art. 2º Revogado.~~\n\n"
        '<p class="Texto_Ementa">Dispõe sobre algo.</p>\n\n'
        '<p class="Fonte_Calibri">Classe pedida à mão.</p>\n'
    )
    resultado = converter_documento_memoria("nota.md", md.encode())
    corpo = resultado["arquivos"][0]["conteudo"]
    obtido = classes_por_texto(corpo)

    assert obtido["NOTA TÉCNICA"] == "Texto_Centralizado_Maiusculas_Negrito"
    assert obtido["Art. 1º Texto citado."] == "Citação"
    assert obtido["codigo  *literal*  indentado"] == "Texto_Mono_Espaçado"
    # O bloco de código não lê marcação e preserva a indentação
    assert "codigo  *literal*<br />  indentado" in corpo
    assert obtido["Texto com trecho riscado no meio."] == "Texto_Justificado"
    assert "<s>trecho riscado</s>" in corpo
    assert obtido["Art. 2º Revogado."] == "Tachado"
    assert obtido["Dispõe sobre algo."] == "Texto_Ementa"
    assert obtido["Classe pedida à mão."] == "Fonte_Calibri"
    assert resultado["avisos"] == []


def test_titulo_numerado_do_markdown_continua_item():
    """Só o título sem número vira título do documento; "1. ASSUNTO" segue numerado pelo SEI."""
    corpo = converter_documento_memoria("nota.md", b"# 1. ASSUNTO\n\nTexto.\n")["arquivos"][0]["conteudo"]
    assert classes_por_texto(corpo)["ASSUNTO"] == "Item_Nivel1"


@pytest.mark.parametrize(
    ("trecho", "aviso"),
    [
        ('<p class="Texto_Justificado">```</p>', "marca de bloco de código"),
    ],
)
def test_validador_aponta_marcacao_markdown_que_sobrou(trecho, aviso):
    from conversorsei.validador import validar_html_sei

    assert any(aviso in falha for falha in validar_html_sei(trecho))


def test_texto_extraido_com_til_nao_vira_riscado():
    """Um "~~" que é conteúdo do documento (ODT, PDF) chega escapado e não risca nada."""
    from conversorsei.md_converter import converter_texto_md_para_blocos, escapar_markdown

    blocos = converter_texto_md_para_blocos(escapar_markdown("valor ~~aproximado~~ do índice"))
    assert "<s>" not in blocos[0]
    assert "~~aproximado~~" in blocos[0]


CONTENT_XML_CLASSES = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:automatic-styles>
  <style:style style:name="P1" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center"/>
   <style:text-properties fo:font-weight="bold" fo:text-transform="uppercase"/>
  </style:style>
  <style:style style:name="P2" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="8cm" fo:text-align="justify"/>
  </style:style>
  <style:style style:name="P3" style:family="paragraph">
   <style:text-properties style:text-line-through-style="solid"/>
  </style:style>
  <style:style style:name="P4" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="4cm"/>
   <style:text-properties fo:font-size="10pt"/>
  </style:style>
  <style:style style:name="P5" style:family="paragraph" style:parent-style-name="Texto_5f_Citação"/>
  <style:style style:name="T1" style:family="text">
   <style:text-properties style:text-line-through-style="solid"/>
  </style:style>
 </office:automatic-styles>
 <office:body>
  <office:text>
   <text:p text:style-name="P1">Portaria de teste</text:p>
   <text:p text:style-name="P2">Dispõe sobre o teste.</text:p>
   <text:p text:style-name="P3">Art. 2º Revogado.</text:p>
   <text:p text:style-name="P4">Art. 1º Texto citado.</text:p>
   <text:p text:style-name="P5">Classe pelo nome do estilo.</text:p>
   <text:p>Texto com <text:span text:style-name="T1">trecho riscado</text:span> no meio.</text:p>
   <text:p>Parágrafo comum.</text:p>
  </office:text>
 </office:body>
</office:document-content>
"""

# O recuo não escolhe classe: "Dispõe..." (8 cm) e "Art. 1º..." (4 cm, 10 pt) seguem texto comum
ESPERADO_ODT = {
    "Portaria de teste": "Texto_Centralizado_Maiusculas_Negrito",
    "Dispõe sobre o teste.": "Texto_Justificado",
    "Art. 2º Revogado.": "Tachado",
    "Art. 1º Texto citado.": "Texto_Justificado",
    "Classe pelo nome do estilo.": "Texto_Citação",
    "Texto com trecho riscado no meio.": "Texto_Justificado",
    "Parágrafo comum.": "Texto_Justificado",
}


def test_odt_leva_caixa_alta_riscado_e_estilo_para_a_classe(tmp_path: Path):
    caminho = tmp_path / "classes.odt"
    with zipfile.ZipFile(caminho, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", CONTENT_XML_CLASSES)

    resultado = converter_documento_memoria("classes.odt", caminho.read_bytes())
    corpo = resultado["arquivos"][0]["conteudo"]

    assert classes_por_texto(corpo) == ESPERADO_ODT, extrair_markdown_odt(caminho)
    assert "<s>trecho riscado</s>" in corpo


def test_paragrafo_recuado_em_fonte_normal_nao_vira_citacao():
    """Recuo sozinho não basta: sem fonte menor o parágrafo segue como texto comum."""
    doc = docx.Document()
    _paragrafo(doc, "Parágrafo recuado por engano.", size=Pt(12)).paragraph_format.left_indent = Cm(3)
    saida = io.BytesIO()
    doc.save(saida)

    corpo = converter_documento_memoria("recuo.docx", saida.getvalue())["arquivos"][0]["conteudo"]
    assert classes_por_texto(corpo)["Parágrafo recuado por engano."] == "Texto_Justificado"


# -- Regressões da revisão do commit 2eb2acc --------------------------------------------


def _corpo(nome: str, conteudo: bytes) -> str:
    resultado = converter_documento_memoria(nome, conteudo, so_corpo=True)
    assert resultado["sucesso"], resultado["erros"]
    return "\n".join(a["conteudo"] for a in resultado["arquivos"])


def test_titulo_numerado_em_caixa_alta_continua_item(tmp_path: Path):
    """"1. ASSUNTO" em caixa alta é item do SEI, e não Texto_Justificado_Maiusculas."""
    xml = CONTENT_XML_CLASSES.split("<office:text>")[0] + (
        "<office:text><text:p>1. ASSUNTO</text:p><text:p>Texto do assunto, com detalhe.</text:p>"
        "</office:text></office:body></office:document-content>"
    )
    caminho = tmp_path / "t.odt"
    with zipfile.ZipFile(caminho, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", xml)

    obtido = classes_por_texto(_corpo("t.odt", caminho.read_bytes()))
    assert obtido["ASSUNTO"] == "Item_Nivel1"


@pytest.mark.parametrize(
    "linha",
    ["```", "```codigo inline```", "~~~~~~~~~~"],
)
def test_cerca_solta_ou_separador_nao_engole_o_documento(linha):
    corpo = _corpo("t.md", f"Antes.\n\n{linha}\n\n# 1. ASSUNTO\n\nDepois.\n".encode())
    obtido = classes_por_texto(corpo)
    assert obtido["ASSUNTO"] == "Item_Nivel1"
    assert obtido["Depois."] == "Texto_Justificado"
    assert "Texto_Mono_Espaçado" not in corpo


@pytest.mark.parametrize(
    ("classe", "esperada"),
    [("Strong", "Texto_Justificado"), ("Heading 1", "Texto_Justificado"), ("Quote", "Texto_Justificado")],
)
def test_classe_que_nao_e_do_sei_e_ignorada(classe, esperada):
    """Estilo do Word que não é classe do SEI não derruba a conversão nem muda a classe."""
    corpo = _corpo("t.md", f'<p class="{classe}">Texto qualquer.</p>\n'.encode())
    assert classes_por_texto(corpo)["Texto qualquer."] == esperada


@pytest.mark.parametrize("markdown", [b"Texto ~~**revogado**~~ fim.\n", b"Texto **~~revogado~~** fim.\n"])
def test_negrito_dentro_de_riscado_nao_vira_asterisco(markdown):
    corpo = _corpo("t.md", markdown)
    assert "**" not in corpo
    assert "<strong><s>revogado</s></strong>" in corpo


def test_riscado_centralizado_em_caixa_alta_e_tachado_no_odt_como_no_word():
    content = _content_odt(
        '<style:style style:name="P1" style:family="paragraph">'
        '<style:paragraph-properties fo:text-align="center"/>'
        '<style:text-properties style:text-line-through-style="solid"/></style:style>',
        '<text:p text:style-name="P1">TEXTO REVOGADO</text:p>',
    )
    assert classes_por_texto(_corpo("t.odt", _odt_bytes(content)))["TEXTO REVOGADO"] == "Tachado"


def test_odt_trecho_riscado_em_paragrafo_riscado_nao_vira_bloco_de_codigo(tmp_path: Path):
    xml = CONTENT_XML_CLASSES.replace(
        '<text:p text:style-name="P3">Art. 2º Revogado.</text:p>',
        '<text:p text:style-name="P3"><text:span text:style-name="T1">Art. 2º</text:span> Revogado.</text:p>',
    )
    caminho = tmp_path / "t.odt"
    with zipfile.ZipFile(caminho, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", xml)

    obtido = classes_por_texto(_corpo("t.odt", caminho.read_bytes()))
    assert obtido["Art. 2º Revogado."] == "Tachado"
    assert obtido["Parágrafo comum."] == "Texto_Justificado"


def test_odt_negrito_so_nas_pontas_nao_e_paragrafo_em_negrito(tmp_path: Path):
    xml = CONTENT_XML_CLASSES.replace(
        '<text:p text:style-name="P1">Portaria de teste</text:p>',
        '<text:p text:style-name="P6"><text:span text:style-name="T2">ANEXO</text:span> I DO '
        '<text:span text:style-name="T2">EDITAL</text:span></text:p>',
    ).replace(
        "</office:automatic-styles>",
        '<style:style style:name="P6" style:family="paragraph">'
        '<style:paragraph-properties fo:text-align="center"/></style:style>'
        '<style:style style:name="T2" style:family="text">'
        '<style:text-properties fo:font-weight="bold"/></style:style></office:automatic-styles>',
    )
    caminho = tmp_path / "t.odt"
    with zipfile.ZipFile(caminho, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", xml)

    assert classes_por_texto(_corpo("t.odt", caminho.read_bytes()))["ANEXO I DO EDITAL"] == (
        "Texto_Centralizado_Maiusculas"
    )


def _docx(montar) -> bytes:
    doc = docx.Document()
    montar(doc)
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def test_recuo_e_fonte_nao_escolhem_classe():
    """Recuo e fonte não decidem a classe: assinatura recuada, com o Normal justificado,
    e documento em Courier seguem texto comum. Ementa e código vêm pelo nome do estilo."""

    def montar(doc):
        doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for texto in ("Brasília, 1 de janeiro de 2026.", "Fulano de Tal"):
            doc.add_paragraph(texto).paragraph_format.left_indent = Cm(8)
        _paragrafo(doc, "Texto citado recuado.", size=Pt(10)).paragraph_format.left_indent = Cm(4)
        _paragrafo(doc, "1. passo = executar()", WD_ALIGN_PARAGRAPH.LEFT, name="Consolas")
        _paragrafo(doc, "Parágrafo em Courier.", name="Courier New")

    obtido = classes_por_texto(_corpo("t.docx", _docx(montar)))
    assert obtido["Brasília, 1 de janeiro de 2026."] == "Texto_Justificado"
    assert obtido["Fulano de Tal"] == "Texto_Justificado"
    assert obtido["Texto citado recuado."] == "Texto_Justificado"
    assert obtido["Parágrafo em Courier."] == "Texto_Justificado"
    assert "Texto_Mono_Espaçado" not in obtido.values()


def test_estilo_explicito_com_numero_nao_conta_como_item():
    """Um "1." num parágrafo de estilo explícito não faz a lista do Word abrir no nível 2."""

    def montar(doc):
        _estilo(doc, "Citação")
        doc.add_paragraph("1. Texto citado que começa com número.", style="Citação")
        p = doc.add_paragraph("Primeiro item da lista")
        _numerar(p, 99)

    obtido = classes_por_texto(_corpo("t.docx", _docx(montar)))
    assert obtido["Primeiro item da lista"] == "Item_Nivel1"


def test_bloco_de_codigo_preserva_indentacao_da_primeira_linha():
    corpo = _corpo("t.md", b"```\n    x = lista[0]\n    y = 1\n```\n")
    assert '<p class="Texto_Mono_Espaçado">    x = lista[0]<br />    y = 1</p>' in corpo


def test_maior_que_no_inicio_de_celula_nao_gera_aviso():
    resultado = converter_documento_memoria("t.md", b"| Idade | Regra |\n|---|---|\n| > 18 anos | Adulto |\n")
    assert resultado["avisos"] == []


def test_marcas_escapadas_no_markdown_nao_geram_aviso():
    """Em .md, ">" e "~~" que chegam ao HTML foram escapados pelo autor: são texto."""
    resultado = converter_documento_memoria("t.md", b"\\> literal\n\n\\~\\~x\\~\\~ literal\n")
    obtido = classes_por_texto(resultado["arquivos"][0]["conteudo"])
    assert obtido == {"> literal": "Texto_Justificado", "~~x~~ literal": "Texto_Justificado"}
    assert resultado["avisos"] == []


# -- Regressões da revisão do commit a177686 --------------------------------------------


def _odt_bytes(content: str, styles: str | None = None) -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
        if styles:
            z.writestr("styles.xml", styles)
    return saida.getvalue()


NS_ODF = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"'
)


def _content_odt(estilos: str, corpo: str) -> str:
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {NS_ODF}>'
        f"<office:automatic-styles>{estilos}</office:automatic-styles>"
        f"<office:body><office:text>{corpo}</office:text></office:body></office:document-content>"
    )


def test_cerca_com_atributos_abre_e_fecha_o_bloco():
    md = "``` python\ncodigo1\n```\n\nTexto normal entre blocos.\n\n```js {x}\ncodigo2\n```\n"
    obtido = classes_por_texto(_corpo("t.md", md.encode()))
    assert obtido == {
        "codigo1": "Texto_Mono_Espaçado",
        "Texto normal entre blocos.": "Texto_Justificado",
        "codigo2": "Texto_Mono_Espaçado",
    }


def test_cerca_de_quatro_crases_aceita_tres_crases_no_conteudo():
    corpo = _corpo("t.md", b"````\nexemplo:\n```\ncodigo\n```\n````\n\nDepois.\n")
    assert classes_por_texto(corpo)["Depois."] == "Texto_Justificado"
    assert "exemplo:<br />```<br />codigo<br />```" in corpo


@pytest.mark.parametrize(
    ("nome", "conteudo"),
    [
        ("t.md", b"~~1. ASSUNTO REVOGADO~~\n\n2. OUTRO\n"),
        (
            "t.odt",
            _odt_bytes(
                _content_odt(
                    '<style:style style:name="P1" style:family="paragraph">'
                    '<style:text-properties style:text-line-through-style="solid"/></style:style>',
                    '<text:p text:style-name="P1">1. ASSUNTO REVOGADO</text:p><text:p>2. OUTRO</text:p>',
                )
            ),
        ),
    ],
)
def test_item_numerado_riscado_continua_na_numeracao(nome, conteudo):
    """Um item revogado continua item: sem isso, o seguinte seria numerado 1 em vez de 2."""
    corpo = _corpo(nome, conteudo)
    obtido = classes_por_texto(corpo)
    assert obtido["ASSUNTO REVOGADO"] == "Item_Nivel1"
    assert obtido["OUTRO"] == "Item_Nivel1"
    assert "<s>" in corpo
    assert "1." not in corpo


def test_odt_titulo_e_item_de_lista_riscados_mantem_o_riscado():
    content = _content_odt(
        '<style:style style:name="P1" style:family="paragraph">'
        '<style:text-properties style:text-line-through-style="solid"/></style:style>'
        '<style:style style:name="T1" style:family="text">'
        '<style:text-properties style:text-line-through-style="solid"/></style:style>',
        '<text:h text:outline-level="2" text:style-name="P1">Seção revogada</text:h>'
        '<text:list><text:list-item><text:p text:style-name="P1">item <text:span text:style-name="T1">todo</text:span>'
        " revogado</text:p></text:list-item></text:list>",
    )
    corpo = _corpo("t.odt", _odt_bytes(content))
    # O título riscado sem número é parágrafo revogado: a classe Tachado já risca
    assert classes_por_texto(corpo)["Seção revogada"] == "Tachado"
    assert "<s>item todo revogado</s>" in corpo


@pytest.mark.parametrize("markdown", [b"Texto ***forte*** fim.\n", b"***forte*** no comeco.\n"])
def test_negrito_com_italico_nao_deixa_asterisco(markdown):
    corpo = _corpo("t.md", markdown)
    assert "*" not in corpo
    assert "<strong><em>forte</em></strong>" in corpo


# -- Regressões da revisão do branch feat/classes-sei -----------------------------------


def test_bloco_de_codigo_mantem_separador_e_link_literais():
    corpo = _corpo("t.md", b"```\n---\nveja [aqui](https://x.com)\n```\n")
    assert '<p class="Texto_Mono_Espaçado">---<br />veja [aqui](https://x.com)</p>' in corpo


def test_bloco_de_codigo_nao_gera_aviso_de_marcacao():
    resultado = converter_documento_memoria("t.md", b"```\n> npm install\n~~x~~\n```\n")
    assert resultado["avisos"] == []


def test_citacao_aninhada_perde_todos_os_marcadores():
    obtido = classes_por_texto(_corpo("t.md", b"> > aninhada\n\n>> colada\n"))
    assert obtido == {"aninhada": "Citação", "colada": "Citação"}


def test_tachado_do_word_sai_sem_s_redundante():
    """O Word e os outros formatos produzem o mesmo HTML para o parágrafo revogado."""
    assert '<p class="Tachado"><strong>a</strong></p>' in _corpo("t.md", b"**~~a~~**\n")


def test_odt_classe_explicita_por_estilo_nomeado_derivado():
    styles = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {NS_ODF}><office:styles>'
        '<style:style style:name="Ementa_20_do_20_despacho" style:family="paragraph"'
        ' style:parent-style-name="Texto_5f_Ementa"/>'
        "</office:styles></office:document-styles>"
    )
    content = _content_odt(
        '<style:style style:name="P1" style:family="paragraph" style:parent-style-name="Ementa_20_do_20_despacho"/>',
        '<text:p text:style-name="P1">Dispõe sobre o teste.</text:p>',
    )
    assert classes_por_texto(_corpo("t.odt", _odt_bytes(content, styles)))["Dispõe sobre o teste."] == "Texto_Ementa"




# -- Regressões da terceira revisão do branch -------------------------------------------


def test_txt_nao_interpreta_citacao_codigo_nem_riscado():
    """.txt não é Markdown: ">", "```" e "~~" são texto, e não geram aviso."""
    resultado = converter_documento_memoria("t.txt", b"> 18 anos ou mais\n\n~~x~~ e ```\n\n~~~\n")
    corpo = resultado["arquivos"][0]["conteudo"]
    obtido = classes_por_texto(corpo)
    assert obtido == {
        "> 18 anos ou mais": "Texto_Justificado",
        "~~x~~ e ```": "Texto_Justificado",
        "~~~": "Texto_Justificado",
    }
    assert resultado["avisos"] == []


def test_paragrafo_riscado_com_link_nao_deixa_til_e_mantem_numero():
    corpo = _corpo("t.md", b"- ~~item com [link](https://x.com)~~\n\n# ~~1. ASSUNTO [ver](https://x.com)~~\n")
    assert "~~" not in corpo
    obtido = classes_por_texto(corpo)
    assert obtido["ASSUNTO ver"] == "Item_Nivel1"
    assert "<s>link</s>" in corpo


def test_odt_styles_xml_malformado_nao_derruba_a_conversao():
    content = _content_odt("", "<text:p>Texto comum.</text:p>")
    corpo = _corpo("t.odt", _odt_bytes(content, "<office:document-styles><quebrado>"))
    assert classes_por_texto(corpo)["Texto comum."] == "Texto_Justificado"


def test_odt_centralizado_herdado_de_estilo_nomeado():
    styles = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {NS_ODF}><office:styles>'
        '<style:style style:name="Titulo_20_centralizado" style:family="paragraph">'
        '<style:paragraph-properties fo:text-align="center"/></style:style>'
        "</office:styles></office:document-styles>"
    )
    content = _content_odt(
        '<style:style style:name="P1" style:family="paragraph" style:parent-style-name="Titulo_20_centralizado"/>',
        '<text:p text:style-name="P1">MINISTÉRIO DA INTEGRAÇÃO</text:p>',
    )
    obtido = classes_por_texto(_corpo("t.odt", _odt_bytes(content, styles)))
    assert obtido["MINISTÉRIO DA INTEGRAÇÃO"] == "Texto_Centralizado_Maiusculas"


def test_estilo_de_corpo_em_celula_de_tabela_nao_vale():
    """Célula com estilo "Texto_Justificado" continua com classe de tabela."""

    def montar(doc):
        _estilo(doc, "Texto_Justificado")
        tabela = doc.add_table(rows=2, cols=1)
        tabela.cell(0, 0).text = "Cabeçalho"
        p = tabela.cell(1, 0).paragraphs[0]
        p.style = doc.styles["Texto_Justificado"]
        p.add_run("Conteúdo da célula")

    obtido = classes_por_texto(_corpo("t.docx", _docx(montar)))
    assert obtido["Conteúdo da célula"] == "Tabela_Texto_Alinhado_Esquerda"


# -- Regressão da quarta revisão ---------------------------------------------------------


def test_titulo_sem_numero_do_odt_vira_item_como_o_titulo_do_word():
    """O Título N do LibreOffice é seção do documento, como o Título N do Word: Item_NivelN."""
    content = _content_odt(
        "",
        '<text:h text:outline-level="1">Introdução</text:h><text:p>Texto.</text:p>'
        '<text:h text:outline-level="2">Contexto</text:h><text:p>Mais texto.</text:p>'
        '<text:h text:outline-level="1">Análise</text:h><text:p>Outro texto.</text:p>',
    )
    obtido = classes_por_texto(_corpo("t.odt", _odt_bytes(content)))
    assert obtido["Introdução"] == "Item_Nivel1"
    assert obtido["Contexto"] == "Item_Nivel2"
    assert obtido["Análise"] == "Item_Nivel1"

    def montar(doc):
        doc.add_heading("Introdução", 1)
        doc.add_paragraph("Texto.")
        doc.add_heading("Contexto", 2)

    word = classes_por_texto(_corpo("t.docx", _docx(montar)))
    assert (word["Introdução"], word["Contexto"]) == (obtido["Introdução"], obtido["Contexto"])


def test_titulo_de_nivel_1_do_markdown_continua_em_caixa_alta():
    assert classes_por_texto(_corpo("t.md", "# Nota técnica\n\nTexto.\n".encode()))["Nota técnica"] == (
        "Texto_Centralizado_Maiusculas_Negrito"
    )


# -- Itens 2 a 7 da quarta revisão -------------------------------------------------------


def test_odt_com_styles_xml_danificado_no_pacote_converte():
    """styles.xml com CRC errado é ignorado: o texto está no content.xml."""
    marcador = b"ESTILOS-DANIFICADOS" * 4
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", _content_odt("", "<text:p>Texto comum.</text:p>"))
        z.writestr("styles.xml", marcador, compress_type=zipfile.ZIP_STORED)
    dados = bytearray(saida.getvalue())
    posicao = dados.index(marcador)
    dados[posicao] ^= 0xFF

    assert classes_por_texto(_corpo("t.odt", bytes(dados)))["Texto comum."] == "Texto_Justificado"


def test_nivel_pedido_pelo_estilo_segue_os_titulos_do_word():
    """--max-nivel limita número digitado. Nível pedido pelo estilo, como o título do Word, fica."""

    def montar(doc):
        _estilo(doc, "Item_Nivel1")
        _estilo(doc, "Item_Nivel2")
        _estilo(doc, "Item_Nivel3")
        doc.add_paragraph("Seção", style="Item_Nivel1")
        doc.add_paragraph("Subseção", style="Item_Nivel2")
        doc.add_paragraph("1.1.1 Detalhe", style="Item_Nivel3")
        doc.add_heading("Título três", level=3)

    resultado = converter_documento_memoria("t.docx", _docx(montar), so_corpo=True, max_nivel=2)
    obtido = classes_por_texto(resultado["arquivos"][0]["conteudo"])
    assert obtido == {
        "Seção": "Item_Nivel1",
        "Subseção": "Item_Nivel2",
        "Detalhe": "Item_Nivel3",
        "Título três": "Item_Nivel3",
    }


def test_fonte_pequena_em_tabela_segue_o_alinhamento():
    """No corpo, célula à direita mantém o alinhamento; o cabeçalho é sempre centralizado."""

    def montar(doc):
        tabela = doc.add_table(rows=2, cols=1)
        _celula(tabela.cell(0, 0), "Valor", WD_ALIGN_PARAGRAPH.RIGHT, 9)
        _celula(tabela.cell(1, 0), "12,5", WD_ALIGN_PARAGRAPH.RIGHT, 8)

    obtido = classes_por_texto(_corpo("t.docx", _docx(montar)))
    assert obtido == {"Valor": "Tabela_Texto_10", "12,5": "Tabela_Texto_Alinhado_Direita"}


def test_link_e_separador_em_bloco_de_codigo_nao_geram_aviso():
    resultado = converter_documento_memoria("t.md", b"```\n---\n[x](https://a.gov.br)\n```\n")
    assert resultado["avisos"] == []


def test_riscado_desligado_no_word_nao_sai_riscado():
    """<w:strike w:val="0"/> é o riscado desmarcado, como o Word grava ao desligá-lo."""

    def montar(doc):
        p = doc.add_paragraph()
        run = p.add_run("texto sem riscado")
        desligado = OxmlElement("w:strike")
        desligado.set(qn("w:val"), "0")
        run._r.get_or_add_rPr().append(desligado)

    corpo = _corpo("t.docx", _docx(montar))
    assert "<s>" not in corpo
    assert classes_por_texto(corpo)["texto sem riscado"] == "Texto_Justificado"


# -- Revisão do branch fix/revisao-4 -----------------------------------------------------


def test_rotulo_de_link_riscado_sai_riscado_e_sem_til():
    corpo = _corpo("t.md", b"Veja o [~~Decreto revogado~~](https://www.planalto.gov.br/x).\n")
    assert "~~" not in corpo
    assert "<s>Decreto revogado</s>" in corpo


def test_estilo_de_caractere_riscado_risca_o_trecho():
    def montar(doc):
        estilo = doc.styles.add_style("Riscado", WD_STYLE_TYPE.CHARACTER)
        estilo.font.strike = True
        p = doc.add_paragraph("Texto com ")
        p.add_run("trecho revogado", style="Riscado")
        p.add_run(" no meio.")

    assert "<s>trecho revogado</s>" in _corpo("t.docx", _docx(montar))


def test_odt_com_content_xml_danificado_tem_mensagem_clara():
    marcador = "<text:p>CONTEUDO-DANIFICADO</text:p>" * 4
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", _content_odt("", marcador), compress_type=zipfile.ZIP_STORED)
    dados = bytearray(saida.getvalue())
    dados[dados.index(marcador.encode())] ^= 0xFF

    resultado = converter_documento_memoria("contrato.odt", bytes(dados))
    assert not resultado["sucesso"]
    assert any("contrato.odt não é um ODT válido" in e for e in resultado["erros"])


def test_odt_styles_xml_com_compressao_nao_suportada_e_ignorado(monkeypatch):
    """Qualquer falha ao ler styles.xml (deflate64, cifrado) só o torna ausente."""
    ler_original = zipfile.ZipFile.read

    def ler(self, nome, *args, **kwargs):
        if nome == "styles.xml":
            raise NotImplementedError("compression type 9 (deflate64)")
        return ler_original(self, nome, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", ler)
    content = _content_odt("", "<text:p>Texto comum.</text:p>")
    assert classes_por_texto(_corpo("t.odt", _odt_bytes(content, "<ok/>")))["Texto comum."] == "Texto_Justificado"


def test_odt_inexistente_sobe_com_a_causa_real(tmp_path: Path):
    """Arquivo que não existe não é "ODT inválido": o erro precisa dizer a causa."""
    with pytest.raises(FileNotFoundError):
        extrair_markdown_odt(tmp_path / "nao_existe.odt")


def test_odt_com_content_xml_malformado_tem_mensagem_clara():
    resultado = converter_documento_memoria("contrato.odt", _odt_bytes("<office:document-content><quebrado"))
    assert any("contrato.odt não é um ODT válido" in e for e in resultado["erros"])


def test_rotulo_riscado_so_de_espaco_mostra_a_url():
    corpo = _corpo("t.md", b"Veja [~~ ~~](https://a.gov.br/x).\n")
    assert "<s>https://a.gov.br/x</s>" in corpo



@pytest.mark.parametrize(
    ("propriedades", "esperado"),
    [
        ('fo:font-weight="bold" style:text-line-through-style="solid"', "<strong><s>forte</s></strong>"),
        ('fo:font-weight="bold" fo:font-style="italic"', "<strong><em>forte</em></strong>"),
    ],
)
def test_odt_marcas_aninhadas_no_trecho_nao_deixam_asterisco(propriedades, esperado):
    """O leitor de ODT escreve "~~**x**~~" e "***x***"; nenhum asterisco pode sobrar."""
    content = _content_odt(
        f'<style:style style:name="T1" style:family="text"><style:text-properties {propriedades}/></style:style>',
        '<text:p>Texto <text:span text:style-name="T1">forte</text:span> fim.</text:p>',
    )
    corpo = _corpo("t.odt", _odt_bytes(content))
    assert "*" not in corpo and "~~" not in corpo
    assert esperado in corpo


# -- Listas do ODT -----------------------------------------------------------------------

ESTILO_DE_LISTA_ODT = (
    '<text:list-style style:name="L1">'
    '<text:list-level-style-number text:level="1" style:num-suffix="." style:num-format="1"/>'
    '<text:list-level-style-number text:level="2" style:num-suffix=")" style:num-format="a"/>'
    '<text:list-level-style-bullet text:level="3" text:bullet-char="•"/>'
    "</text:list-style>"
)


def test_lista_numerada_do_odt_escreve_o_numero_no_texto():
    """Como a <ol> do HTML: o número fica no texto e não entra na numeração do SEI."""
    content = _content_odt(
        ESTILO_DE_LISTA_ODT,
        '<text:list text:style-name="L1">'
        "<text:list-item><text:p>Primeira pergunta</text:p>"
        "<text:list><text:list-item><text:p>detalhe um</text:p></text:list-item>"
        "<text:list-item><text:p>detalhe dois</text:p>"
        "<text:list><text:list-item><text:p>marcador</text:p></text:list-item></text:list>"
        "</text:list-item></text:list></text:list-item>"
        '<text:list-item text:start-value="5"><text:p>Quinta pergunta</text:p></text:list-item>'
        "</text:list>",
    )
    corpo = _corpo("t.odt", _odt_bytes(content))
    obtido = classes_por_texto(corpo)
    assert obtido["1. Primeira pergunta"] == "Texto_Justificado"
    assert obtido["a) detalhe um"] == "Texto_Justificado"
    assert obtido["b) detalhe dois"] == "Texto_Justificado"
    assert obtido["5. Quinta pergunta"] == "Texto_Justificado"
    assert 'Texto_Justificado">marcador</li>' in corpo
    assert "Item_Nivel" not in corpo


def test_lista_com_marcadores_do_odt_mantem_os_niveis():
    content = _content_odt(
        "",
        "<text:list><text:list-item><text:p>um</text:p>"
        "<text:list><text:list-item><text:p>um.a</text:p></text:list-item></text:list>"
        "</text:list-item><text:list-item><text:p>dois</text:p></text:list-item></text:list>",
    )
    corpo = _corpo("t.odt", _odt_bytes(content))
    assert re.sub(r"\s", "", corpo) == (
        '<ul><liclass="Texto_Justificado">um<ul><liclass="Texto_Justificado">um.a</li></ul></li>'
        '<liclass="Texto_Justificado">dois</li></ul>'
    )
