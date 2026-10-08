import docx
from docx.oxml.ns import qn

from conversorsei.md_converter import converter_md_para_blocos, converter_md_para_html, markdown_para_docx
from conversorsei.web import converter_documento_memoria


def test_md_conversao_tabela_e_links(tmp_path):
    md_file = tmp_path / "documento.md"
    md_content = """# 1. ASSUNTO

1.1. Texto referente à [Lei nº 12.587/2012](https://www.planalto.gov.br/ccivil_03/_ato2011-2014/2012/lei/l12587.htm).

## 2. ANÁLISE

2.1. Avaliação do impacto:

| Parâmetro | Situação |
| --- | --- |
| Impacto Financeiro | Baixo |
| Complexidade | Moderada |

<red>**ALERTA:** Este é um documento confidencial.</red>
"""
    md_file.write_text(md_content, encoding="utf-8")

    # Gera DOCX intermediário
    docx_file = tmp_path / "documento.docx"
    markdown_para_docx(md_file, docx_file)
    assert docx_file.is_file()

    # Converte para blocos SEI
    blocos = converter_md_para_blocos(md_file)
    assert len(blocos) >= 4

    html = converter_md_para_html(md_file)
    assert 'class="Item_Nivel1"' in html
    assert 'class="Item_Nivel2"' in html
    assert "https://www.planalto.gov.br/ccivil_03/_ato2011-2014/2012/lei/l12587.htm" in html
    assert "ccivil<em>" not in html
    assert "color:#C00000" in html
    assert "<table" in html


def test_md_escapes_de_barra_invertida(tmp_path):
    """Markdown gerado por modelo vem com "1\\." e "\\---": a barra não pode virar texto.

    Com a barra intacta, o título deixa de ser reconhecido como item de nível 1 e os
    parágrafos "1.1." ficam órfãos, fazendo o contador do SEI começar em "0.1.".
    """
    md_file = tmp_path / "documento.md"
    md_file.write_text(
        "\\---\n\n"
        "## 1\\. ASSUNTO\n\n"
        "1.1. Trata-se de análise do Projeto de Lei nº 7\\.210, de 2025.\n\n"
        "\\---\n\n"
        "## 2\\. REFERÊNCIAS\n\n"
        "2.1. [Lei nº 12.587](https://www.planalto.gov.br/l12587.htm), que institui a PNMU.\n",
        encoding="utf-8",
    )

    html = converter_md_para_html(md_file)
    assert "\\" not in html
    # Título vira item de nível 1; a linha separadora escapada é descartada
    assert html.count('class="Item_Nivel1"') == 2
    assert html.count('class="Item_Nivel2"') == 2
    assert ">ASSUNTO<" in html
    assert "Projeto de Lei nº 7.210" in html


def test_md_escape_de_asterisco_nao_vira_enfase(tmp_path):
    """Um "\\*" escapado é asterisco literal, não abertura de ênfase."""
    md_file = tmp_path / "documento.md"
    md_file.write_text("O campo \\*obrigatório\\* deve ser preenchido.\n", encoding="utf-8")

    html = converter_md_para_html(md_file)
    assert "*obrigatório*" in html
    assert "<em>" not in html


def test_travessao_e_meia_risca_ficam_no_texto(tmp_path):
    """O texto sai como veio, como no .docx: a meia-risca entre números é intervalo.

    A limpeza antiga trocava o travessão por vírgula e apagava a meia-risca, e
    "2020–2023" virava "20202023".
    """
    md_file = tmp_path / "documento.md"
    md_file.write_text(
        "Produtos Industrializados – IPI –, na aquisição. Vigência 2020–2023, p. 10–15. O prazo — curto — vence.\n",
        encoding="utf-8",
    )

    html = converter_md_para_html(md_file)
    assert "Industrializados – IPI –, na aquisição" in html
    assert "2020–2023" in html
    assert "p. 10–15" in html
    assert "O prazo — curto — vence" in html


def test_md_tabela_com_linha_incompleta_formata_todas_as_celulas(tmp_path):
    """Linha do Markdown com menos colunas que a maior linha da tabela.

    A tabela do DOCX nasce com a largura da maior linha, então as células que a linha
    curta não preencheu existem mesmo assim. Elas precisam da mesma margem e, no
    cabeçalho, da mesma tarja cinza, sob pena de sair uma célula destoante no SEI.
    """
    md_file = tmp_path / "irregular.md"
    md_file.write_text(
        "| Parâmetro | Situação |\n"
        "| --- | --- |\n"
        "| Impacto |\n"
        "| Complexidade | Moderada |\n",
        encoding="utf-8",
    )

    docx_file = tmp_path / "irregular.docx"
    markdown_para_docx(md_file, docx_file)

    tabela = docx.Document(str(docx_file)).tables[0]
    assert len(tabela.columns) == 2

    for linha in tabela.rows:
        for celula in linha.cells:
            assert celula._tc.find(qn("w:tcPr")).find(qn("w:tcMar")) is not None

    for celula in tabela.rows[0].cells:
        shd = celula._tc.find(qn("w:tcPr")).find(qn("w:shd"))
        assert shd is not None and shd.get(qn("w:fill")) == "E6E6E6"

    # A célula que o Markdown não preencheu fica vazia, e não com o texto da coluna ao lado
    assert tabela.cell(1, 0).text == "Impacto"
    assert tabela.cell(1, 1).text == ""


def test_md_com_fim_de_linha_do_windows_nao_ganha_quebra_extra(tmp_path):
    """CRLF não pode virar quebra a mais dentro de um bloco de várias linhas.

    A leitura passou a ser binária (para a interface web entregar bytes), então o \\r
    deixou de ser removido pelo modo texto e precisa ser normalizado no parser.
    """
    md_file = tmp_path / "windows.md"
    md_file.write_bytes(b"# 1. ASSUNTO\r\n\r\n<red>linha um\r\nlinha dois</red>\r\n")

    blocos = converter_md_para_blocos(md_file)
    vermelho = next(b for b in blocos if "C00000" in b)

    assert "\r" not in vermelho
    assert vermelho.count("<br />") == 1


def test_titulo_do_markdown_sai_sem_cor():
    """Regressão: o DOCX intermediário dava cor aos títulos (azul em # e ##, cinza em ###),
    e o HTML levava um <span style="color:..."> que o SEI mostrava. A aparência do título
    vem da classe, e não de uma cor escolhida pelo conversor."""
    md = "# NOTA INFORMATIVA\n\n## 1. ASSUNTO\n\n1.1. Texto.\n\n### 1.2. Detalhe\n\nMais texto.\n"
    res = converter_documento_memoria("nota.md", md.encode("utf-8"), so_corpo=True)
    html = res["arquivos"][0]["conteudo"]
    assert "ASSUNTO" in html and "Detalhe" in html
    assert "color:" not in html


def test_tabela_do_markdown_sai_centralizada():
    """A tabela do Markdown passa pelo mesmo conversor da do Word e sai centralizada."""
    md = "Texto.\n\n| Norma | Objeto |\n|---|---|\n| NBR 10151 | Ruído ambiental |\n"
    html = converter_documento_memoria("t.md", md.encode("utf-8"), so_corpo=True)["arquivos"][0]["conteudo"]
    # Markdown não informa largura: a tabela ocupa a página
    assert '<table align="center" border="1" cellpadding="1" cellspacing="1" width="100%">' in html


def test_md_link_com_parenteses_no_endereco_e_sem_cor_embutida(tmp_path):
    """O link vai inteiro, e sai como o do Word: sem <span> de cor nem <u> na saída."""
    md_file = tmp_path / "links.md"
    md_file.write_text("Ver [a lei](https://pt.wikipedia.org/wiki/Lei_(direito)) hoje.\n", encoding="utf-8")

    html = converter_md_para_html(md_file)
    assert '<a href="https://pt.wikipedia.org/wiki/Lei_(direito)">a lei</a> hoje.' in html
    assert "color" not in html and "<u>" not in html


def test_md_link_com_esquema_inseguro_fica_so_o_texto(tmp_path):
    md_file = tmp_path / "links.md"
    md_file.write_text("Veja [clique](javascript:alert(1)) agora.\n", encoding="utf-8")

    html = converter_md_para_html(md_file)
    assert "<a" not in html
    assert "Veja clique agora." in html


def test_md_marcador_mais_vira_item_de_lista(tmp_path):
    """"+ " é marcador no CommonMark, e a página já trata esse texto colado como Markdown."""
    md_file = tmp_path / "lista.md"
    md_file.write_text("+ primeiro\n+ segundo\n", encoding="utf-8")

    html = converter_md_para_html(md_file)
    assert '<li class="Texto_Justificado">primeiro</li>' in html
    assert "+" not in html


def test_txt_mantem_a_barra_invertida_do_texto():
    """No .txt a barra é caractere, como no texto extraído de HTML, ODT e PDF."""
    from conversorsei.core import converter_bytes

    texto = "Pasta C:\\dados\\(2026), regex \\d+\\. e ~~nada~~ riscado.\n> resposta citada\n"
    res = converter_bytes("nota.txt", texto.encode())
    assert res.sucesso, res.erros
    html = res.arquivos[0].conteudo
    assert "Pasta C:\\dados\\(2026), regex \\d+\\. e ~~nada~~ riscado." in html
    assert "&gt; resposta citada" in html
