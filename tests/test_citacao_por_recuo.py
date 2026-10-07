"""Citação pelo recuo e pela fonte menor, que só vale a pedido.

Documento real quase nunca usa estilo: a citação vem recuada e em fonte menor, e sem a
regra sairia como texto comum. A regra adivinha pela medida, então fica desligada por
padrão. A conversão conta os parágrafos com forma de citação, e quem converte decide
(o botão da interface web, ou --citacao-por-recuo na CLI).
"""
from __future__ import annotations

import html
import io
import json
import re
import zipfile

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from conversorsei.cli import main
from conversorsei.web import converter_documento_memoria

RE_PARAGRAFO = re.compile(r'<p class="([^"]+)">(.*?)</p>', re.DOTALL)
CITACAO = "Art. 13. O titular do serviço divulgará os dados."


def classes_por_texto(corpo: str) -> dict[str, str]:
    return {html.unescape(re.sub(r"<[^>]+>", "", t)).strip(): c for c, t in RE_PARAGRAFO.findall(corpo)}


def converter(nome: str, conteudo: bytes, aplicar: bool = False) -> tuple[dict[str, str], int]:
    res = converter_documento_memoria(nome, conteudo, so_corpo=True, citacao_por_recuo=aplicar)
    assert res["sucesso"], res["erros"]
    corpo = "\n".join(a["conteudo"] for a in res["arquivos"])
    return classes_por_texto(corpo), res["citacoes_por_recuo"]


def _paragrafo(doc, texto: str, tamanho: float | None = 12, recuo_cm: float = 0, alinhamento=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if alinhamento is None else alinhamento
    run = p.add_run(texto)
    if tamanho is not None:
        run.font.size = Pt(tamanho)
    if recuo_cm:
        p.paragraph_format.left_indent = Cm(recuo_cm)
    return p


def _salvar(doc) -> bytes:
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _corpo_comum(doc, tamanho: float | None = 12) -> None:
    _paragrafo(doc, "Trata-se de análise da proposição, com os fundamentos a seguir.", tamanho)
    _paragrafo(doc, "A lei dispõe sobre a transparência dos dados do transporte coletivo.", tamanho)


def docx_com_citacao() -> bytes:
    doc = docx.Document()
    _corpo_comum(doc)
    _paragrafo(doc, CITACAO, tamanho=10, recuo_cm=4)
    _paragrafo(doc, "(...)", tamanho=10, recuo_cm=4)
    _paragrafo(doc, "Nesse sentido, está sendo desenvolvida uma plataforma.", 12)
    return _salvar(doc)


def test_citacao_por_recuo_e_contada_mas_nao_aplicada_por_padrao():
    classes, encontradas = converter("nt.docx", docx_com_citacao())
    assert encontradas == 2
    assert classes[CITACAO] == "Texto_Justificado"
    assert classes["(...)"] == "Texto_Justificado"


def test_resultado_traz_o_inicio_de_cada_citacao():
    """A interface mostra o começo dos parágrafos, e quem decide não precisa caçá-los na prévia."""
    longo = "Art. 14. A concessionária publicará os horários e as tarifas de cada linha no portal."
    doc = docx.Document()
    # Corpo maior que as citações, para a fonte menor não virar a predominante
    _corpo_comum(doc)
    _corpo_comum(doc)
    _paragrafo(doc, CITACAO, tamanho=10, recuo_cm=4)
    _paragrafo(doc, longo, tamanho=10, recuo_cm=4)
    res = converter_documento_memoria("nt.docx", _salvar(doc), so_corpo=True)

    assert res["citacoes_trechos"] == [CITACAO, " ".join(longo.split()[:12]) + "…"]


@pytest.mark.parametrize("aplicar", [False, True])
def test_odt_traz_o_inicio_de_cada_citacao(aplicar):
    res = converter_documento_memoria("nt.odt", odt_com_citacao(), so_corpo=True, citacao_por_recuo=aplicar)
    assert len(res["citacoes_trechos"]) == res["citacoes_por_recuo"] > 0
    assert all(trecho and "~~" not in trecho for trecho in res["citacoes_trechos"])


def test_citacao_por_recuo_aplicada_a_pedido():
    classes, encontradas = converter("nt.docx", docx_com_citacao(), aplicar=True)
    assert encontradas == 2
    assert classes[CITACAO] == "Citação"
    assert classes["(...)"] == "Citação"
    # O texto comum em volta não muda
    assert classes["Nesse sentido, está sendo desenvolvida uma plataforma."] == "Texto_Justificado"


def test_documento_todo_em_fonte_pequena_nao_tem_citacao():
    """A fonte é comparada com a do corpo: recuado na mesma fonte do texto não é citação."""
    doc = docx.Document()
    _corpo_comum(doc, tamanho=10)
    _paragrafo(doc, CITACAO, tamanho=10, recuo_cm=4)
    classes, encontradas = converter("t.docx", _salvar(doc), aplicar=True)
    assert encontradas == 0
    assert classes[CITACAO] == "Texto_Justificado"


@pytest.mark.parametrize(
    "montar, texto, esperada",
    [
        # Assinatura recuada na fonte do texto
        (lambda d: _paragrafo(d, "Fulano de Tal", 12, 8), "Fulano de Tal", "Texto_Justificado"),
        # Recuo pequeno demais
        (lambda d: _paragrafo(d, CITACAO, 10, 1), CITACAO, "Texto_Justificado"),
        # Legenda centralizada
        (
            lambda d: _paragrafo(d, "Figura 1 - Mapa", 10, 4, WD_ALIGN_PARAGRAPH.CENTER),
            "Figura 1 - Mapa",
            "Texto_Centralizado",
        ),
        # Item numerado digitado continua item, como sem a regra
        (lambda d: _paragrafo(d, "3.1. Texto do item recuado.", 10, 4), "Texto do item recuado.", "Item_Nivel2"),
    ],
)
def test_parecido_com_citacao_mas_nao_e(montar, texto, esperada):
    doc = docx.Document()
    _corpo_comum(doc)
    montar(doc)
    classes, encontradas = converter("t.docx", _salvar(doc), aplicar=True)
    assert encontradas == 0
    assert classes[texto] == esperada


def test_celula_de_tabela_nao_conta_como_citacao():
    doc = docx.Document()
    _corpo_comum(doc)
    celula = doc.add_table(rows=1, cols=1).cell(0, 0)
    run = celula.paragraphs[0].add_run("Valor na tabela")
    run.font.size = Pt(10)
    celula.paragraphs[0].paragraph_format.left_indent = Cm(4)
    _, encontradas = converter("t.docx", _salvar(doc), aplicar=True)
    assert encontradas == 0


def test_recuo_gravado_como_start_e_fonte_do_padrao_do_documento():
    """O Word atual grava w:start, e o corpo costuma deixar a fonte no padrão (docDefaults)."""
    doc = docx.Document()
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for antigo in rpr.findall(qn("w:sz")):
        rpr.remove(antigo)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    doc.styles["Normal"].font.size = None
    _corpo_comum(doc, tamanho=None)
    p = _paragrafo(doc, CITACAO, tamanho=10)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:start"), str(int(4 * 1440 / 2.54)))
    p._p.get_or_add_pPr().append(ind)

    classes, encontradas = converter("t.docx", _salvar(doc), aplicar=True)
    assert encontradas == 1
    assert classes[CITACAO] == "Citação"


CONTENT_ODT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:automatic-styles>
  <style:style style:name="P1" style:family="paragraph">
   <style:paragraph-properties fo:text-align="justify"/>
  </style:style>
  <style:style style:name="P2" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="1.5748in" fo:text-align="justify"/>
   <style:text-properties fo:font-size="10pt"/>
  </style:style>
  <style:style style:name="P3" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="8cm"/>
  </style:style>
 </office:automatic-styles>
 <office:body>
  <office:text>
   <text:p text:style-name="P1">Trata-se de análise da proposição, com os fundamentos a seguir.</text:p>
   <text:p text:style-name="P1">A lei dispõe sobre a transparência dos dados do transporte coletivo.</text:p>
   <text:p text:style-name="P2">Art. 13. O titular do serviço divulgará os dados.</text:p>
   <text:p text:style-name="P2">3.1. Item recuado em fonte menor.</text:p>
   <text:p text:style-name="P3">Fulano de Tal</text:p>
  </office:text>
 </office:body>
</office:document-content>
"""

STYLES_ODT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0">
 <office:styles>
  <style:default-style style:family="paragraph">
   <style:text-properties fo:font-size="12pt"/>
  </style:default-style>
 </office:styles>
</office:document-styles>
"""


def odt_com_citacao() -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", CONTENT_ODT)
        z.writestr("styles.xml", STYLES_ODT)
    return saida.getvalue()


@pytest.mark.parametrize("aplicar, esperada", [(False, "Texto_Justificado"), (True, "Citação")])
def test_odt_segue_a_mesma_regra_do_word(aplicar, esperada):
    """Recuo em polegadas e fonte do estilo padrão (styles.xml) valem como no Word."""
    classes, encontradas = converter("nt.odt", odt_com_citacao(), aplicar=aplicar)
    assert encontradas == 1
    assert classes[CITACAO] == esperada
    assert classes["Item recuado em fonte menor."] == "Item_Nivel2"
    assert classes["Fulano de Tal"] == "Texto_Justificado"


def test_formatos_sem_recuo_nao_contam():
    _, encontradas = converter("t.md", "Texto.\n\n> Citação do Markdown.\n".encode(), aplicar=True)
    assert encontradas == 0


def test_cli_avisa_e_aplica_com_a_flag(tmp_path, capsys):
    entrada = tmp_path / "nt.docx"
    entrada.write_bytes(docx_com_citacao())

    assert main([str(entrada), "-o", str(tmp_path / "a")]) == 0
    saida = capsys.readouterr().out
    assert "2 parágrafos recuados e em fonte menor que a do texto parecem citação" in saida
    assert "--citacao-por-recuo" in saida
    assert 'class="Citação"' not in (tmp_path / "a" / "nt_SEI.html").read_text(encoding="utf-8")

    assert main([str(entrada), "-o", str(tmp_path / "b"), "--citacao-por-recuo"]) == 0
    assert "foram convertidos como Citação" in capsys.readouterr().out
    assert (tmp_path / "b" / "nt_SEI.html").read_text(encoding="utf-8").count('class="Citação"') == 2

    assert main([str(entrada), "-o", str(tmp_path / "c"), "--json"]) == 0
    assert json.loads(capsys.readouterr().out)[0]["citacoes_por_recuo"] == 2


def test_cli_e_web_geram_o_mesmo_html_com_a_opcao(tmp_path):
    entrada = tmp_path / "nt.docx"
    entrada.write_bytes(docx_com_citacao())
    main([str(entrada), "-o", str(tmp_path), "--citacao-por-recuo"])
    web = converter_documento_memoria("nt.docx", entrada.read_bytes(), citacao_por_recuo=True)
    assert web["arquivos"][0]["conteudo"] == (tmp_path / "nt_SEI.html").read_text(encoding="utf-8")


# Citação entre aspas: o parágrafo inteiro entre aspas que passa de uma linha segue o mesmo
# pedido da citação pelo recuo, e vale em todos os formatos
ENTRE_ASPAS = (
    "“O titular do serviço divulgará, em formato aberto, os dados de oferta e de demanda "
    "do transporte coletivo, atualizados a cada mês.”"
)
CURTA = "“O titular divulgará os dados.”"


def converter_aspas(nome: str, conteudo: bytes, aplicar: bool = False, **opcoes) -> tuple[dict[str, str], dict]:
    res = converter_documento_memoria(nome, conteudo, so_corpo=True, citacao_por_recuo=aplicar, **opcoes)
    assert res["sucesso"], res["erros"]
    corpo = "\n".join(a["conteudo"] for a in res["arquivos"])
    return classes_por_texto(corpo), res


def docx_com_paragrafos(*textos: str, alinhamento=None) -> bytes:
    doc = docx.Document()
    _corpo_comum(doc)
    for texto in textos:
        _paragrafo(doc, texto, alinhamento=alinhamento)
    return _salvar(doc)


@pytest.mark.parametrize("aplicar, esperada", [(False, "Texto_Justificado"), (True, "Citação")])
def test_paragrafo_entre_aspas_com_mais_de_uma_linha_e_sugerido(aplicar, esperada):
    classes, res = converter_aspas("nt.docx", docx_com_paragrafos(ENTRE_ASPAS), aplicar)
    assert res["citacoes_entre_aspas"] == 1
    assert res["citacoes_por_recuo"] == 0
    assert res["citacoes_trechos"] == [" ".join(ENTRE_ASPAS.split()[:12]) + "…"]
    # As aspas ficam no texto: a regra só troca a classe
    assert classes[ENTRE_ASPAS] == esperada


@pytest.mark.parametrize(
    "texto",
    [
        ENTRE_ASPAS.removesuffix("”") + ".” (BRASIL, 2020, p. 3).",
        '"' + ENTRE_ASPAS[1:-1] + '"',
        "«" + ENTRE_ASPAS[1:-1] + "»",
        # Quebra de linha dentro das aspas já é mais de uma linha
        "“Art. 13. O titular divulgará os dados.\nParágrafo único. A divulgação será mensal.”",
    ],
)
def test_variantes_da_citacao_entre_aspas(texto):
    doc = docx.Document()
    _corpo_comum(doc)
    p = _paragrafo(doc, "")
    for i, linha in enumerate(texto.split("\n")):
        if i:
            p.runs[-1].add_break()
        p.add_run(linha)
    _, res = converter_aspas("t.docx", _salvar(doc), aplicar=True)
    assert res["citacoes_entre_aspas"] == 1


@pytest.mark.parametrize(
    "texto",
    [
        # Cabe numa linha: fala curta, e não citação em parágrafo próprio
        CURTA,
        # Aspas fechadas no meio: dois trechos citados dentro do texto
        "“Dados abertos” e “transparência ativa” são os dois conceitos que a lei usa ao tratar da divulgação.",
        # Texto depois das aspas que não é referência
        ENTRE_ASPAS + " Assim, a obrigação já existe e só falta regulamentar o prazo.",
        # Aspas que não abrem o parágrafo
        "Diz a lei: " + ENTRE_ASPAS,
    ],
)
def test_entre_aspas_mas_nao_e_citacao(texto):
    classes, res = converter_aspas("t.docx", docx_com_paragrafos(texto), aplicar=True)
    assert res["citacoes_entre_aspas"] == 0
    assert classes[texto] == "Texto_Justificado"


def test_entre_aspas_centralizado_ou_item_numerado_segue_a_regra_de_antes():
    classes, res = converter_aspas(
        "t.docx", docx_com_paragrafos(ENTRE_ASPAS, alinhamento=WD_ALIGN_PARAGRAPH.CENTER), aplicar=True
    )
    assert res["citacoes_entre_aspas"] == 0
    assert classes[ENTRE_ASPAS] == "Texto_Centralizado"

    classes, res = converter_aspas("t.docx", docx_com_paragrafos("3.1. " + ENTRE_ASPAS), aplicar=True)
    assert res["citacoes_entre_aspas"] == 0
    assert classes[ENTRE_ASPAS] == "Item_Nivel2"


def test_recuo_e_aspas_contam_separados_e_os_trechos_seguem_o_documento():
    doc = docx.Document()
    _corpo_comum(doc)
    _corpo_comum(doc)
    _paragrafo(doc, ENTRE_ASPAS)
    _paragrafo(doc, CITACAO, tamanho=10, recuo_cm=4)
    # Recuado, em fonte menor e entre aspas: conta uma vez só, pelo recuo
    _paragrafo(doc, ENTRE_ASPAS.replace("O titular", "A concessionária"), tamanho=10, recuo_cm=4)
    _, res = converter_aspas("nt.docx", _salvar(doc))
    assert (res["citacoes_por_recuo"], res["citacoes_entre_aspas"]) == (2, 1)
    assert [t.split()[0] for t in res["citacoes_trechos"]] == ["“O", "Art.", "“A"]


@pytest.mark.parametrize(
    "nome, conteudo, opcoes",
    [
        ("t.md", f"Texto do parecer.\n\n{ENTRE_ASPAS}\n", {}),
        ("t.txt", f"Texto do parecer.\n\n{ENTRE_ASPAS}\n", {}),
        ("t.html", f"<p>Texto do parecer.</p><p>{ENTRE_ASPAS}</p>", {}),
        ("t.pdf", b"", {"texto_reconhecido": f"Texto do parecer.\n\n{ENTRE_ASPAS}\n"}),
    ],
)
def test_aspas_valem_nos_formatos_sem_recuo(nome, conteudo, opcoes):
    dados = conteudo.encode() if isinstance(conteudo, str) else conteudo
    classes, res = converter_aspas(nome, dados, aplicar=True, **opcoes)
    assert res["citacoes_entre_aspas"] == 1
    assert classes[ENTRE_ASPAS] == "Citação"


def test_citacao_do_markdown_continua_explicita_e_nao_e_sugerida():
    classes, res = converter_aspas("t.md", f"Texto.\n\n> {ENTRE_ASPAS}\n".encode())
    assert res["citacoes_entre_aspas"] == 0
    assert classes[ENTRE_ASPAS] == "Citação"


def odt_com_paragrafos(*corpos: str) -> bytes:
    paragrafos = "".join(f'<text:p text:style-name="{estilo}">{texto}</text:p>' for estilo, texto in corpos)
    content = CONTENT_ODT.split("<office:text>")[0] + f"<office:text>{paragrafos}</office:text>" + (
        CONTENT_ODT.split("</office:text>")[1]
    )
    content = content.replace(
        "</office:automatic-styles>",
        '<style:style style:name="T1" style:family="text"><style:text-properties fo:font-style="italic"/>'
        "</style:style></office:automatic-styles>",
    )
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
        z.writestr("styles.xml", STYLES_ODT)
    return saida.getvalue()


def test_odt_em_italico_e_recuado_entre_aspas_contam_uma_vez():
    """O asterisco do itálico não esconde as aspas, e o ODT não conta o parágrafo duas vezes."""
    recuada = ENTRE_ASPAS.replace("O titular", "A concessionária")
    conteudo = odt_com_paragrafos(
        ("P1", "Trata-se de análise da proposição, com os fundamentos a seguir."),
        ("P1", "A lei dispõe sobre a transparência dos dados do transporte coletivo."),
        ("P1", f'<text:span text:style-name="T1">{ENTRE_ASPAS}</text:span>'),
        ("P2", recuada),
    )
    classes, res = converter_aspas("nt.odt", conteudo)
    assert (res["citacoes_por_recuo"], res["citacoes_entre_aspas"]) == (1, 1)
    assert len(res["citacoes_trechos"]) == 2
    assert not any("*" in t for t in res["citacoes_trechos"])

    classes, res = converter_aspas("nt.odt", conteudo, aplicar=True)
    assert classes[ENTRE_ASPAS] == "Citação"
    assert classes[recuada] == "Citação"


def test_cli_explica_a_citacao_entre_aspas(tmp_path, capsys):
    entrada = tmp_path / "nt.md"
    entrada.write_text(f"Texto do parecer.\n\n{ENTRE_ASPAS}\n", encoding="utf-8")

    assert main([str(entrada), "-o", str(tmp_path / "a")]) == 0
    assert "1 parágrafo entre aspas e com mais de uma linha parece citação" in capsys.readouterr().out

    assert main([str(entrada), "-o", str(tmp_path / "b"), "--json"]) == 0
    assert json.loads(capsys.readouterr().out)[0]["citacoes_entre_aspas"] == 1
