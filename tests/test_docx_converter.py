import io
import re

import docx
import pytest
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from conversorsei.docx_converter import (
    RE_STRIP_ITEM,
    converter_docx_para_blocos,
    converter_docx_para_html,
    remover_tags_vazias,
)


def test_docx_strip_numeracao_item():
    exemplos = [
        ('<span style="color:#003366"><strong>1. ASSUNTO</strong></span>', "ASSUNTO"),
        ('<span style="color:#333333"><strong>3.4.1. Análise Técnica</strong></span>', "Análise Técnica"),
        ("<strong>2.1. Referência</strong>", "Referência"),
        # Só a numeração em negrito: a tag de fechamento fica entre o número e o texto
        ("<strong>1.</strong> ASSUNTO", "ASSUNTO"),
        ("<strong>1.1.</strong> Descrição", "Descrição"),
        ('<span style="color:#003366"><strong>3.4.1.</strong></span> Análise', "Análise"),
    ]
    for bruto, esperado in exemplos:
        limpo = RE_STRIP_ITEM.sub(r"\1\2", bruto, count=1)
        assert esperado in limpo
        assert "1. ASSUNTO" not in limpo
        assert "3.4.1." not in limpo
        # A numeração sai inteira: o SEI renumera sozinho e o resto viraria "1. 1. ASSUNTO"
        assert not re.search(r">\s*\d", limpo)
        # E sem deixar para trás as tags que envolviam só o número
        assert not re.search(r"<(\w+)[^>]*></\1>", remover_tags_vazias(limpo))


def test_docx_conversao_completa(tmp_path):
    doc_path = tmp_path / "minuta.docx"
    doc = docx.Document()

    # Adiciona seções e parágrafos
    doc.add_heading("1. ASSUNTO", level=1)
    doc.add_paragraph("1.1. Primeiro parágrafo de detalhamento.")
    doc.add_paragraph("1.2. Segundo parágrafo com texto justificado.")

    # Adiciona tabela
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Item"
    tbl.cell(0, 1).text = "Descrição"
    tbl.cell(1, 0).text = "1"
    tbl.cell(1, 1).text = "Texto do item"

    doc.save(doc_path)

    blocos = converter_docx_para_blocos(doc_path)
    assert len(blocos) >= 3

    html_corpo = converter_docx_para_html(doc_path)
    assert 'class="Item_Nivel1"' in html_corpo
    assert 'class="Item_Nivel2"' in html_corpo
    assert "<table" in html_corpo
    assert 'class="Tabela_Texto_Centralizado"' in html_corpo

    # Determinismo
    html_2 = converter_docx_para_html(doc_path)
    assert html_corpo == html_2


def test_docx_preserva_formatacao_em_quebra_de_linha(tmp_path):
    """Um run com <w:br> mantém negrito e ordem do texto (regressão)."""
    doc_path = tmp_path / "quebra.docx"
    doc = docx.Document()
    p = doc.add_paragraph()
    r = p.add_run("Primeira linha")
    r.bold = True
    r.add_break()
    r.add_text("Segunda linha")
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert "<strong>Primeira linha<br />Segunda linha</strong>" in html_corpo


def test_docx_preserva_texto_visivel_em_revisoes_controles_e_campos(tmp_path):
    """Texto exibido pelo Word não pode sumir por estar dentro de um contêiner XML."""
    doc = docx.Document()
    p = doc.add_paragraph("Antes ")

    def run_com_texto(texto):
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = texto
        run.append(t)
        return run

    insercao = OxmlElement("w:ins")
    insercao.append(run_com_texto("INSERIDO"))
    p._p.append(insercao)

    controle = OxmlElement("w:sdt")
    conteudo = OxmlElement("w:sdtContent")
    conteudo.append(run_com_texto("CONTROLE"))
    controle.append(conteudo)
    p._p.append(controle)

    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "DATE")
    campo.append(run_com_texto("DATA VISÍVEL"))
    p._p.append(campo)

    exclusao = OxmlElement("w:del")
    exclusao.append(run_com_texto("EXCLUÍDO"))
    p._p.append(exclusao)
    p.add_run(" Depois")

    arquivo = tmp_path / "revisoes.docx"
    doc.save(arquivo)
    html = converter_docx_para_html(arquivo)
    assert "Antes INSERIDOCONTROLEDATA VISÍVEL Depois" in html
    assert "EXCLUÍDO" not in html


def test_docx_preserva_paragrafo_dentro_de_controle_de_conteudo(tmp_path):
    """Um controle no corpo pode envolver o parágrafo inteiro."""
    doc = docx.Document()
    paragrafo = doc.add_paragraph("Texto do controle")
    controle = OxmlElement("w:sdt")
    conteudo = OxmlElement("w:sdtContent")
    doc.element.body.remove(paragrafo._p)
    conteudo.append(paragrafo._p)
    controle.append(conteudo)
    doc.element.body.insert(0, controle)

    arquivo = tmp_path / "controle.docx"
    doc.save(arquivo)
    assert "Texto do controle" in converter_docx_para_html(arquivo)


def _adicionar_numeracao(p, ilvl: int = 0, num_id: int = 99) -> None:
    """Marca o parágrafo como item de lista numerada do Word (w:numPr).

    O numId 99 não existe no numbering.xml do template, então o conversor usa
    o formato decimal padrão em vez do marcador do numId 1.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    el_ilvl = OxmlElement("w:ilvl")
    el_ilvl.set(qn("w:val"), str(ilvl))
    el_num = OxmlElement("w:numId")
    el_num.set(qn("w:val"), str(num_id))
    numPr.append(el_ilvl)
    numPr.append(el_num)
    pPr.append(numPr)


def test_docx_lista_numerada_sem_titulo_comeca_no_nivel1(tmp_path):
    """Sem títulos no documento, a lista numerada do Word abre em Item_Nivel1."""
    doc_path = tmp_path / "lista.docx"
    doc = docx.Document()
    _adicionar_numeracao(doc.add_paragraph("Primeiro item da lista"))
    _adicionar_numeracao(doc.add_paragraph("Segundo item da lista"))
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert 'class="Item_Nivel1"' in html_corpo
    assert 'class="Item_Nivel2"' not in html_corpo


def test_docx_lista_numerada_apos_titulo_vira_subitem(tmp_path):
    """Havendo título antes, a lista numerada passa a ser subitem (Item_Nivel2)."""
    doc_path = tmp_path / "titulo_e_lista.docx"
    doc = docx.Document()
    doc.add_heading("1. ASSUNTO", level=1)
    _adicionar_numeracao(doc.add_paragraph("Item subordinado ao título"))
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert 'class="Item_Nivel1"' in html_corpo
    assert 'class="Item_Nivel2"' in html_corpo


def test_docx_lista_com_marcadores_recebe_classe_sei(tmp_path):
    doc_path = tmp_path / "bullets.docx"
    doc = docx.Document()
    doc.add_paragraph("Primeiro marcador", style="List Bullet")
    doc.add_paragraph("Segundo marcador", style="List Bullet")
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert "<ul>" in html_corpo
    assert '<li class="Texto_Justificado">Primeiro marcador</li>' in html_corpo
    assert "<li>" not in html_corpo


def test_docx_tabela_com_gridspan_e_vmerge(tmp_path):
    """Células mescladas viram colspan/rowspan e não duplicam conteúdo."""
    from docx.oxml.ns import qn

    doc_path = tmp_path / "merge.docx"
    doc = docx.Document()
    tbl = doc.add_table(rows=3, cols=2)
    tbl.cell(0, 0).text = "Cabecalho unico"
    tbl.cell(0, 1).text = "x"
    tbl.cell(0, 0).merge(tbl.cell(0, 1))          # horizontal -> colspan=2
    tbl.cell(1, 0).text = "Lateral"
    tbl.cell(1, 0).merge(tbl.cell(2, 0))          # vertical -> rowspan=2
    tbl.cell(1, 1).text = "Direita cima"
    tbl.cell(2, 1).text = "Direita baixo"
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert 'colspan="2"' in html_corpo
    assert 'rowspan="2"' in html_corpo
    # A célula mesclada aparece uma vez só
    assert html_corpo.count("Lateral") == 1
    assert doc.tables[0]._tbl.findall(qn("w:tr"))  # sanidade: a tabela existe no XML


def test_docx_imagem_vira_base64(tmp_path):
    """Imagem do documento entra embutida em Base64, sem referência a arquivo externo."""
    import base64

    # PNG 1x1 transparente
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    img_path = tmp_path / "ponto.png"
    img_path.write_bytes(png)

    doc_path = tmp_path / "com_imagem.docx"
    doc = docx.Document()
    doc.add_paragraph("Texto antes da imagem.")
    doc.add_picture(str(img_path))
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert 'src="data:image/png;base64,' in html_corpo
    assert "ponto.png" not in html_corpo
    assert 'style="max-width:100%;"' in html_corpo


def test_docx_hyperlink_vira_ancora(tmp_path):
    doc_path = tmp_path / "link.docx"
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("Consulte o ")
    # Hyperlink real, com relacionamento externo
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = p.part.relate_to("https://www.gov.br/", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "portal gov.br"
    run.append(texto)
    link.append(run)
    p._p.append(link)
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert '<a href="https://www.gov.br/">portal gov.br</a>' in html_corpo


def test_parece_texto_corrido_distingue_titulo_de_paragrafo():
    from conversorsei.docx_converter import parece_texto_corrido

    # Títulos de seção: curtos, sem pontuação interna, com ou sem ponto final
    assert not parece_texto_corrido("1. ASSUNTO")
    assert not parece_texto_corrido("2. ANÁLISE TÉCNICA")
    assert not parece_texto_corrido("3. Conclusão")
    assert not parece_texto_corrido("4. Conclusão.")
    assert not parece_texto_corrido("2. Sistema de Gestão, Controle e Auditoria")
    # Parágrafos de texto corrido
    assert parece_texto_corrido(
        "1. Trata-se da necessidade de contratação de serviços técnicos especializados "
        "para implantação da plataforma."
    )
    assert parece_texto_corrido("4. Diante do exposto, encaminham-se os autos à Subsecretaria.")
    assert parece_texto_corrido("1. O processo foi analisado pela equipe tecnica.")


def test_docx_paragrafos_numerados_nao_viram_titulo(tmp_path):
    """Despacho com parágrafos numerados não pode sair em caixa alta com tarja cinza."""
    doc_path = tmp_path / "despacho.docx"
    doc = docx.Document()
    doc.add_paragraph("À Subsecretaria de Planejamento, Orçamento e Administração (SPOA)")
    doc.add_paragraph(
        "1. Trata-se da necessidade de contratação de serviços técnicos especializados "
        "para implantação e sustentação da infraestrutura tecnológica."
    )
    doc.add_paragraph(
        "2. Com o objetivo de dar continuidade aos procedimentos necessários à contratação, "
        "solicita-se a designação da respectiva equipe de planejamento."
    )
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert html_corpo.count('class="Paragrafo_Numerado_Nivel1"') == 2
    assert "Item_Nivel1" not in html_corpo
    # A numeração digitada sai, quem numera é o contador do SEI
    assert ">Trata-se da necessidade" in html_corpo


def test_docx_titulos_curtos_continuam_como_secao(tmp_path):
    """Nota técnica com seções curtas mantém Item_Nivel1 (caixa alta e tarja são desejadas)."""
    doc_path = tmp_path / "nota.docx"
    doc = docx.Document()
    doc.add_paragraph("1. ASSUNTO")
    doc.add_paragraph("1.1. Texto do assunto tratado no expediente.")
    doc.add_paragraph("2. ANÁLISE")
    doc.add_paragraph("2.1. Texto da análise.")
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert html_corpo.count('class="Item_Nivel1"') == 2
    assert html_corpo.count('class="Item_Nivel2"') == 2
    assert "Paragrafo_Numerado" not in html_corpo


def _numerar_pelo_word(p, num_id=5, ilvl=0):
    """Aplica numeração automática decimal ao parágrafo (numId 5 do template padrão)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numPr = OxmlElement("w:numPr")
    elem_ilvl = OxmlElement("w:ilvl")
    elem_ilvl.set(qn("w:val"), str(ilvl))
    numPr.append(elem_ilvl)
    elem_num = OxmlElement("w:numId")
    elem_num.set(qn("w:val"), str(num_id))
    numPr.append(elem_num)
    p._p.get_or_add_pPr().append(numPr)


def test_docx_numeracao_automatica_de_paragrafos_nao_vira_titulo(tmp_path):
    """Despacho numerado pela barra do Word: o número não está em p.text, mas o texto é corrido."""
    doc_path = tmp_path / "despacho_auto.docx"
    doc = docx.Document()
    doc.add_paragraph("À Subsecretaria de Planejamento, Orçamento e Administração (SPOA)")
    for texto in (
        "Trata-se da necessidade de contratação de serviços técnicos especializados para "
        "implantação e sustentação da infraestrutura tecnológica.",
        "Com o objetivo de dar continuidade aos procedimentos necessários à contratação, "
        "solicita-se a designação da respectiva equipe de planejamento.",
    ):
        _numerar_pelo_word(doc.add_paragraph(texto))
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert html_corpo.count('class="Paragrafo_Numerado_Nivel1"') == 2
    assert "Item_Nivel1" not in html_corpo


def test_docx_secoes_com_numeracao_automatica_seguem_titulo(tmp_path):
    """Seções curtas numeradas pelo Word continuam em Item_Nivel, com a tarja cinza."""
    doc_path = tmp_path / "nota_auto.docx"
    doc = docx.Document()
    for texto in ("ASSUNTO", "REFERÊNCIAS", "ANÁLISE", "CONCLUSÃO"):
        _numerar_pelo_word(doc.add_paragraph(texto))
    doc.save(doc_path)

    html_corpo = converter_docx_para_html(doc_path)
    assert html_corpo.count('class="Item_Nivel1"') == 4
    assert "Paragrafo_Numerado" not in html_corpo


def test_docx_numeracao_so_em_negrito_nao_duplica(tmp_path):
    """Word com "1." em negrito e o texto fora do negrito.

    A numeração fica num run próprio, então no HTML a tag de fechamento entra entre o
    número e o texto. Sem removê-la junto, o número sobrevive e o SEI, que renumera
    pela classe Item_Nivel, mostra "1. 1. ASSUNTO".
    """
    doc_path = tmp_path / "negrito.docx"
    doc = docx.Document()

    p = doc.add_paragraph()
    p.add_run("1.").bold = True
    p.add_run(" ASSUNTO")

    sub = doc.add_paragraph()
    sub.add_run("1.1.").bold = True
    sub.add_run(" Trata-se de análise da minuta encaminhada pela unidade demandante.")

    doc.save(doc_path)

    html = converter_docx_para_html(doc_path)
    assert 'class="Item_Nivel1"' in html
    assert "ASSUNTO" in html
    assert "Trata-se de análise" in html
    # Nenhum número digitado sobra depois de uma tag
    assert not re.search(r">\s*\d+(\.\d+)*\.", html)


def test_tabela_sai_centralizada_mesmo_alinhada_a_esquerda_no_word():
    """Colada no SEI, a tabela ficava à esquerda. Toda tabela sai centralizada e com a
    largura da página, qualquer que seja o alinhamento no Word."""
    doc = docx.Document()
    doc.add_paragraph("Texto antes da tabela.")
    tabela = doc.add_table(rows=2, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabela.cell(0, 0).text = "Norma"
    tabela.cell(1, 0).text = "NBR 10151"
    blocos = converter_docx_para_blocos(_salvar_em_memoria(doc))

    abertura = next(b for b in blocos if b.startswith("<table")).split("\n", 1)[0]
    assert 'align="center"' in abertura
    assert 'width="100%"' in abertura


def test_larguras_da_tabela_vao_como_atributo_e_nao_no_style():
    """Regressão do Safari: ao copiar, ele troca a largura do style pelo valor medido em
    px (width:100% virava 718px), e a tabela saía estreita e à esquerda no SEI. A largura
    em atributo passa intacta, na tabela e nas células."""
    doc = docx.Document()
    tabela = doc.add_table(rows=2, cols=3)
    for i, texto in enumerate(["Aspecto", "Emissão do veículo", "Ruído ambiental"]):
        tabela.cell(0, i).text = texto
    blocos = converter_docx_para_blocos(_salvar_em_memoria(doc))
    html = next(b for b in blocos if b.startswith("<table"))

    assert not re.search(r'style="[^"]*width', html)
    celulas = re.findall(r"<td\b[^>]*>", html)
    assert all('width="' in td for td in celulas[:3])



def _png(largura: int, altura: int) -> bytes:
    """PNG cinza do tamanho pedido, montado sem biblioteca de imagens."""
    import struct
    import zlib

    def pedaco(tipo: bytes, dados: bytes) -> bytes:
        return struct.pack(">I", len(dados)) + tipo + dados + struct.pack(">I", zlib.crc32(tipo + dados))

    cabecalho = struct.pack(">IIBBBBB", largura, altura, 8, 0, 0, 0, 0)
    linhas = b"".join(b"\x00" + b"\x80" * largura for _ in range(altura))
    return (
        b"\x89PNG\r\n\x1a\n" + pedaco(b"IHDR", cabecalho) + pedaco(b"IDAT", zlib.compress(linhas))
        + pedaco(b"IEND", b"")
    )


def _tamanhos_das_imagens(imagens: list[tuple[bytes, float]]) -> list[tuple[int, int]]:
    """Converte um DOCX com cada imagem na largura pedida (em polegadas) e lê width e height."""
    from docx.shared import Inches

    doc = docx.Document()  # Carta, margens de 1,25": área de texto de 6"
    for conteudo, polegadas in imagens:
        doc.add_picture(io.BytesIO(conteudo), width=Inches(polegadas))
    html = "\n".join(converter_docx_para_blocos(_salvar_em_memoria(doc)))
    return [(int(w), int(h)) for w, h in re.findall(r'<img alt="" width="(\d+)" height="(\d+)"', html)]


def test_imagem_sai_com_a_proporcao_da_pagina_do_word():
    """Sem largura no <img>, o SEI mostrava a foto no tamanho original (1.600 px de câmera).

    A largura toda da área de texto vale 800 px no SEI; a imagem guarda a fração da página
    que ocupa no Word, e a altura segue a proporção. Vai como atributo, e não no style.
    """
    foto = _png(1600, 800)
    assert _tamanhos_das_imagens([(foto, 6.0), (foto, 1.5)]) == [(800, 400), (200, 100)]


def test_imagem_pequena_nao_e_ampliada():
    """Uma imagem de 100 px esticada no Word até a largura da página continua com 100 px."""
    assert _tamanhos_das_imagens([(_png(100, 50), 6.0)]) == [(100, 50)]


def test_dimensoes_em_pixels_le_png_gif_e_jpeg():
    from conversorsei.docx_converter import dimensoes_em_pixels

    assert dimensoes_em_pixels(_png(30, 20)) == (30, 20)
    assert dimensoes_em_pixels(b"GIF89a" + (640).to_bytes(2, "little") + (480).to_bytes(2, "little") + b"\x00") == (640, 480)
    # JPEG: SOI, um APP0 qualquer e o SOF0 com altura 915 e largura 1600
    app0 = b"\xff\xe0" + (16).to_bytes(2, "big") + b"JFIF\x00" + b"\x00" * 9
    sof0 = b"\xff\xc0" + (17).to_bytes(2, "big") + b"\x08" + (915).to_bytes(2, "big") + (1600).to_bytes(2, "big")
    assert dimensoes_em_pixels(b"\xff\xd8" + app0 + sof0 + b"\x03" + b"\x00" * 9) == (1600, 915)
    assert dimensoes_em_pixels(b"formato desconhecido") is None


def _salvar_em_memoria(doc) -> bytes:
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _largura_no_html(doc) -> str:
    html = next(b for b in converter_docx_para_blocos(_salvar_em_memoria(doc)) if b.startswith("<table"))
    return re.match(r'<table[^>]*\bwidth="(\d+)%"', html).group(1)


def _definir_tblW(tabela, valor: str, tipo: str) -> None:
    tblW = tabela._tbl.tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), valor)
    tblW.set(qn("w:type"), tipo)


@pytest.mark.parametrize(
    "valor, tipo, esperada",
    [
        ("metade", "dxa", "50"),
        ("2500", "pct", "50"),
        ("80%", "pct", "80"),
        # tblW sem medida: vale a grade, que o python-docx monta na largura da página
        ("0", "auto", "100"),
        # Mais larga que a área de texto: fica na página
        ("dobro", "dxa", "100"),
    ],
)
def test_tabela_mantem_a_proporcao_que_tem_no_word(valor, tipo, esperada):
    """A tabela sai com a largura que tem no Word, em % da área de texto, e centralizada."""
    doc = docx.Document()
    secao = doc.sections[-1]
    area = secao.page_width.twips - secao.left_margin.twips - secao.right_margin.twips
    valor = {"metade": str(area // 2), "dobro": str(area * 2)}.get(valor, valor)
    tabela = doc.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "A"
    _definir_tblW(tabela, valor, tipo)
    assert _largura_no_html(doc) == esperada


def test_tabela_aninhada_ocupa_a_celula_inteira():
    doc = docx.Document()
    externa = doc.add_table(rows=1, cols=1)
    _definir_tblW(externa, "2500", "pct")
    interna = externa.cell(0, 0).add_table(rows=1, cols=1)
    interna.cell(0, 0).text = "dentro"
    _definir_tblW(interna, "1000", "pct")
    html = next(b for b in converter_docx_para_blocos(_salvar_em_memoria(doc)) if b.startswith("<table"))
    assert re.findall(r'<table[^>]*\bwidth="(\d+)%"', html) == ["50", "100"]


def test_docx_numero_em_run_proprio_separado_por_tab_continua_item(tmp_path):
    """Regressão: o run só de tab entre "1." e o título sumia do texto, e "1.ASSUNTO"
    deixava de ser item numerado."""
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("1.").bold = True
    p.add_run("\t")
    p.add_run("ASSUNTO").bold = True
    doc.add_paragraph("Texto do parágrafo.")
    arquivo = tmp_path / "item.docx"
    doc.save(arquivo)

    html = converter_docx_para_html(arquivo)
    item = re.search(r'<p class="Item_Nivel1">(.*?)</p>', html)
    assert item, html
    assert item.group(1) == "<strong>ASSUNTO</strong>"


def test_docx_marcador_com_recuo_padrao_fica_no_primeiro_nivel(tmp_path):
    """O recuo direto de 1,27 cm é o padrão de uma lista de primeiro nível, e não subitem."""
    from docx.shared import Cm

    doc = docx.Document()
    for texto in ("Primeiro", "Segundo"):
        doc.add_paragraph(texto, style="List Bullet").paragraph_format.left_indent = Cm(1.27)
    arquivo = tmp_path / "marcadores.docx"
    doc.save(arquivo)

    html = converter_docx_para_html(arquivo)
    assert "<ul>" in html
    assert "margin-left" not in html


def test_docx_marcador_aninhado_pelo_ilvl_mantem_hierarquia(tmp_path):
    """O nível do marcador vem do ilvl da numeração do Word, sem recuo direto."""
    doc = docx.Document()
    for texto, ilvl in (("Item", 0), ("Subitem", 1), ("Outro item", 0)):
        p = doc.add_paragraph(texto, style="List Bullet")
        _adicionar_numeracao(p, ilvl=ilvl, num_id=1)
    arquivo = tmp_path / "aninhada.docx"
    doc.save(arquivo)

    html = converter_docx_para_html(arquivo)
    assert re.search(r"Item<ul><li[^>]*>Subitem</li></ul></li>", html), html


def test_docx_controle_de_conteudo_dentro_de_celula(tmp_path):
    """O texto de um controle de conteúdo na célula não pode sumir, como já não some no corpo."""
    doc = docx.Document()
    tabela = doc.add_table(rows=2, cols=1)
    tabela.cell(0, 0).text = "Campo"
    celula = tabela.cell(1, 0)
    paragrafo = celula.paragraphs[0]
    paragrafo.add_run("Valor do formulário")
    controle = OxmlElement("w:sdt")
    conteudo = OxmlElement("w:sdtContent")
    celula._tc.remove(paragrafo._p)
    conteudo.append(paragrafo._p)
    controle.append(conteudo)
    celula._tc.append(controle)
    arquivo = tmp_path / "formulario.docx"
    doc.save(arquivo)

    assert "Valor do formulário" in converter_docx_para_html(arquivo)


def test_estilo_do_paragrafo_nao_e_buscado_de_novo_a_cada_leitura(monkeypatch):
    """Regressão de desempenho: o p.style do python-docx percorre a lista de estilos a cada
    leitura, e tomava mais de 90% do tempo de um documento longo (8 s em 3.300 parágrafos,
    com a aba do navegador congelada). O número de buscas não pode crescer com o documento."""
    from docx.styles.styles import Styles

    chamadas = 0
    original = Styles.get_by_id

    def contar(self, *args, **kwargs):
        nonlocal chamadas
        chamadas += 1
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Styles, "get_by_id", contar)
    doc = docx.Document()
    for i in range(300):
        doc.add_paragraph(f"{i}. Item" if i % 10 == 0 else f"Parágrafo {i}.")
    buffer = io.BytesIO()
    doc.save(buffer)
    converter_docx_para_blocos(io.BytesIO(buffer.getvalue()))
    assert chamadas <= 5, f"{chamadas} buscas de estilo para 300 parágrafos"
