"""
md_converter.py — Converte documentos Markdown (.md) em DOCX institucional e/ou blocos HTML SEI.
"""
from __future__ import annotations

import codecs
import copy
import html as htmlmod
import io
import re
from dataclasses import replace
from pathlib import Path

import docx

# docx.Document e a fabrica; o tipo do objeto vive em docx.document
from docx.document import Document as DocumentoWord
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from conversorsei.docx_converter import converter_docx_para_blocos, profundidade_item
from conversorsei.entrada import FonteDocumento, ler_bytes
from conversorsei.formatacao import (
    CLASSE_ASSINATURA,
    CitacaoPorRecuo,
    Formatacao,
    classe_por_formatacao,
    classe_sei_pelo_nome,
    texto_em_maiusculas,
)

RE_INICIO_NUMERADO = re.compile(r"^\d+(?:\.\d+)*\.?\s")
RE_LINK = re.compile(r"\[([^\]]*)\]\(([^)]+)\)")
RE_IMAGEM_MD = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
MARCA_QUEBRA_ODT = "\ue000"
# Célula do ODT que ocupa várias colunas: a largura vai no início do texto, entre dois
# caracteres de uso privado. Só marcar_colunas_mescladas escreve a marca, e só
# colunas_mescladas a lê
RE_COLUNA_MESCLADA_ODT = re.compile(r"^\ue001(\d+)\ue002")
# A nota só conta pela definição ("[^1]: texto" no início da linha): a chamada solta
# confunde-se com texto comum, como a expressão "[^0-9]"
RE_NOTA_MD = re.compile(r"^\s{0,3}\[\^[^\]\s]+\]:")
RE_SEQUENCIA_UTF8 = re.compile(rb"[\xc2-\xdf][\x80-\xbf]|[\xe0-\xef][\x80-\xbf]{2}|[\xf0-\xf4][\x80-\xbf]{3}")
RE_RED = re.compile(r"<red>(.*?)</red>", re.DOTALL)
RE_HTML_P = re.compile(r"^<p\s+class=[\"']([^\"']+)[\"']>(.*?)</p>$", re.IGNORECASE | re.DOTALL)
# "> texto" (citação) e "```" (bloco de código): blocos do Markdown que têm classe própria no SEI
# Todos os níveis de citação, com ou sem espaço entre eles ("> > x" e ">> x")
RE_CITACAO_MD = re.compile(r"^(?:>\s?)+(.*)$")
# Cerca de código como no CommonMark: três ou mais crases e, opcionalmente, a linguagem
# e outros atributos ("``` python", "```js {x}"), sem crase depois delas: "```x```" na
# mesma linha é código inline, e não bloco. Cerca de til (~~~) fica de fora de propósito:
# em .txt, uma linha de tis é separador, e duas delas prenderiam o texto entre elas.
RE_CERCA_CODIGO = re.compile(r"^(`{3,})[^`]*$")
# Riscado: linha inteira vira Tachado, trecho vira <s>
RE_RISCADO = re.compile(r"~~(.+?)~~")
COR_VERMELHA = "C00000"

# Escape de Markdown: barra invertida seguida de pontuação ASCII (CommonMark).
# Modelos de linguagem costumam devolver "1\. ASSUNTO" ou "\---", e a barra não pode
# chegar ao texto final nem esconder a numeração de quem detecta títulos e itens.
RE_ESCAPE_MD = re.compile(r"\\([!-/:-@\[-`{-~])")
# Marca o caractere escapado durante a leitura da marcação inline, para que um "\*"
# não seja lido como ênfase. Restaurado na hora de escrever o run no docx.
SENTINELA = "\x00"
RE_SENTINELA = re.compile(rf"{SENTINELA}(\d+){SENTINELA}")


# O texto extraído de PDF ou ODT atravessa o conversor de Markdown antes de virar
# HTML, então um asterisco ou colchete do documento seria lido como marcação
# ("*cautela*" virava itálico). A barra invertida é o escape que format_paragraph desfaz
# ao escrever o run.
RE_MARCACAO_MD = re.compile(r"([*\[\]`~])")
RE_ABERTURA_DE_BLOCO_MD = re.compile(r"^([#>|+*-])")


def escapar_markdown(texto: str) -> str:
    """Neutraliza a marcação Markdown que vier no texto extraído de um documento."""
    texto = RE_MARCACAO_MD.sub(r"\\\1", texto)
    # No início da linha, estes caracteres abrem título, citação, tabela ou lista
    return RE_ABERTURA_DE_BLOCO_MD.sub(r"\\\1", texto)


def marcar_colunas_mescladas(texto: str, largura: int) -> str:
    """Célula de tabela Markdown que leva junto o número de colunas que ocupa."""
    return f"\ue001{largura}\ue002{texto}" if largura > 1 else texto


def colunas_mescladas(celula: str) -> tuple[int, str]:
    """Desfaz marcar_colunas_mescladas: devolve a largura e o texto sem a marca."""
    m = RE_COLUNA_MESCLADA_ODT.match(celula)
    return (int(m.group(1)), celula[m.end() :]) if m else (1, celula)


def desescapar_markdown(texto: str) -> str:
    r"""Texto como o leitor o vê: sem as barras de escape ("1\." vira "1.").

    A marca de quebra de linha do ODT vira espaço. Este texto serve para classificar e
    para mostrar trechos na tela, onde o caractere de uso privado colaria as palavras.
    """
    return RE_ESCAPE_MD.sub(r"\1", texto).replace(MARCA_QUEBRA_ODT, " ")


def proteger_escapes(texto: str) -> str:
    """Substitui o caractere escapado por uma marca neutra à leitura da marcação inline."""
    return RE_ESCAPE_MD.sub(lambda m: f"{SENTINELA}{ord(m.group(1))}{SENTINELA}", texto)


def restaurar_escapes(texto: str) -> str:
    """Desfaz proteger_escapes(), devolvendo o caractere literal."""
    return RE_SENTINELA.sub(lambda m: chr(int(m.group(1))), texto)


def sanitizar_travessoes(texto: str) -> str:
    """Remove travessões/meia-riscas (— ou –) substituindo por vírgulas ou ajustando pontuação.

    O hífen comum (-) é rigorosamente preservado.
    """
    texto = re.sub(r"\s+[—–]\s+", ", ", texto)
    texto = texto.replace("—", "").replace("–", "")
    texto = re.sub(r",\s*,", ",", texto)
    # O travessão de aposto ("o IPI –, na aquisição") deixa espaço antes da pontuação
    texto = re.sub(r"\s+([,;.])", r"\1", texto)
    return texto


def definir_fundo_celula(cell, fill_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def definir_margens_celula(cell, top: int = 100, bottom: int = 100, left: int = 150, right: int = 150) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tcPr.append(tcMar)


def aplicar_bordas_tabela(table) -> None:
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def adicionar_hiperlink(paragraph, url: str, text: str, riscado: bool = False) -> None:
    """Adiciona um hyperlink real (azul, sublinhado, Calibri 11) ao parágrafo do docx."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if riscado:
        rPr.append(OxmlElement("w:strike"))
    new_run.append(rPr)

    # A quebra de linha do ODT dentro do link vira <w:br/> no próprio run, como no texto comum
    for indice, trecho in enumerate(text.split(MARCA_QUEBRA_ODT)):
        if indice:
            new_run.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = trecho
        new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def parse_markdown_inline(text: str) -> list[tuple]:
    segments = []
    pos = 0
    for m in RE_RED.finditer(text):
        if m.start() > pos:
            segments.extend(_parse_links_emphasis(text[pos : m.start()], None))
        segments.extend(_parse_links_emphasis(m.group(1), COR_VERMELHA))
        pos = m.end()
    if pos < len(text):
        segments.extend(_parse_links_emphasis(text[pos:], None))
    return segments


def _parse_links_emphasis(text: str, color: str | None) -> list[tuple]:
    segments = []
    pos = 0
    for m in RE_LINK.finditer(text):
        if m.start() > pos:
            segments.extend(_split_emphasis(text[pos : m.start()], color))
        rotulo = (m.group(1) or "").strip() or m.group(2)
        segments.append(("link", m.group(2), rotulo))
        pos = m.end()
    if pos < len(text):
        segments.extend(_split_emphasis(text[pos:], color))
    return segments


def _split_emphasis(
    text: str, color: str | None = None, bold: bool = False, italic: bool = False, riscado: bool = False
) -> list[tuple]:
    """Separa negrito, itálico e riscado, aceitando uma marca dentro da outra.

    O ODT escreve um trecho em negrito e riscado como "~~**x**~~". Sem ler o
    interior de cada marca, os asteriscos chegariam ao texto final.
    """
    # "***x***" vem antes de "**": é como o ODT escreve negrito com itálico
    pattern = r"(\*\*\*.+?\*\*\*|\*\*.*?\*\*|~~.*?~~|\*.*?\*)"
    parts = re.split(pattern, text)
    result: list[tuple] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("***") and part.endswith("***") and len(part) >= 7:
            result.extend(_split_emphasis(part[3:-3], color, True, True, riscado))
        elif part.startswith("**") and part.endswith("**") and len(part) >= 4:
            result.extend(_split_emphasis(part[2:-2], color, True, italic, riscado))
        elif part.startswith("~~") and part.endswith("~~") and len(part) >= 5:
            result.extend(_split_emphasis(part[2:-2], color, bold, italic, True))
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            result.extend(_split_emphasis(part[1:-1], color, bold, True, riscado))
        else:
            result.append(("text", part, bold, italic, color, riscado))
    return result


def html_paragraph_to_markdown(line: str) -> tuple[str, str] | None:
    m = RE_HTML_P.match(line.strip())
    if not m:
        return None
    cls, inner = m.group(1), m.group(2)
    inner = re.sub(r"<br\s*/?>", "\n", inner, flags=re.IGNORECASE)
    inner = re.sub(r"<strong>(.*?)</strong>", r"**\1**", inner, flags=re.IGNORECASE | re.DOTALL)
    inner = re.sub(r"<em>(.*?)</em>", r"*\1*", inner, flags=re.IGNORECASE | re.DOTALL)
    inner = re.sub(r"<[^>]+>", "", inner)
    inner = htmlmod.unescape(inner).strip()
    return cls, inner


def iter_red_block_paragraphs(lines: list[str], start_idx: int) -> tuple[list[str], int]:
    coletadas = []
    i = start_idx
    while i < len(lines):
        coletadas.append(lines[i])
        if "</red>" in lines[i]:
            break
        i += 1
    texto = "\n".join(coletadas)
    texto = texto.replace("<red>", "").replace("</red>", "")
    partes = [p.strip() for p in re.split(r"\n\s*\n", texto) if p.strip()]
    return partes, min(i + 1, len(lines))


# Formatação de run já montada, por combinação de atributos. Aplicar a formatação pelas
# propriedades do python-docx custa caro (cada atributo é uma busca no XML), e um documento
# longo tem dezenas de milhares de runs com poucas combinações distintas: a primeira de
# cada combinação é montada pela API, e as seguintes recebem uma cópia do mesmo w:rPr.
_RPR_POR_FORMATO: dict[tuple, object] = {}


def add_text_run(
    paragraph,
    run_text: str,
    bold: bool,
    italic: bool,
    color: str | None,
    is_heading: bool,
    level: int,
    riscado: bool = False,
):
    chave = (bold, italic, color, is_heading, level, riscado)
    partes = restaurar_escapes(run_text).split("\n")
    for idx, parte in enumerate(partes):
        if idx:
            paragraph.add_run().add_break()
        if not parte:
            continue
        r = paragraph.add_run(parte)
        modelo = _RPR_POR_FORMATO.get(chave)
        if modelo is not None:
            r._r.insert(0, copy.deepcopy(modelo))
            continue
        formatar_run(r, bold, italic, color, is_heading, level, riscado)
        _RPR_POR_FORMATO[chave] = copy.deepcopy(r._r.rPr)


def formatar_run(r, bold: bool, italic: bool, color: str | None, is_heading: bool, level: int, riscado: bool) -> None:
    """Formatação de um run de texto do Markdown, pela API do python-docx."""
    r.font.name = "Calibri"
    # Riscado vale também em título: um título revogado continua revogado
    if riscado:
        r.font.strike = True

    if is_heading:
        r.font.bold = True
        if level == 1:
            r.font.size = Pt(13)
        elif level == 2:
            r.font.size = Pt(12)
        else:
            r.font.size = Pt(11)
        # Sem cor padrão: o docx_converter leva a cor do run para o HTML, e o título
        # sairia azul (ou cinza) no SEI, onde quem define a aparência é a classe
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    else:
        r.font.size = Pt(11)
        r.font.bold = bold
        r.font.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(
    p, text: str, is_heading: bool = False, level: int = 1, is_center: bool = False, is_left: bool = False
):
    riscado_inteiro = texto_inteiro_riscado(text)
    if riscado_inteiro:
        # O riscado vale para todos os trechos, links inclusive. Lido como marca inline,
        # "~~a [link](url) b~~" deixaria os "~~" no texto: os links são separados antes
        text = text.replace("~~", "")
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2)

    if is_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif is_left:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif not is_heading:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    clean_text = proteger_escapes(sanitizar_travessoes(text))

    for seg in parse_markdown_inline(clean_text):
        if seg[0] == "link":
            _, url, rotulo = seg
            # O rótulo não passa pela leitura de ênfase: "[~~Decreto~~](url)" deixaria os
            # "~~" no texto do link. Rótulo todo riscado vira link riscado
            rotulo_riscado = texto_inteiro_riscado(rotulo)
            if rotulo_riscado:
                # Sem as marcas, um rótulo só de espaço deixaria o link invisível: vale a URL
                rotulo = rotulo.replace("~~", "").strip() or url
            adicionar_hiperlink(
                p, restaurar_escapes(url), restaurar_escapes(rotulo), riscado_inteiro or rotulo_riscado
            )
            continue

        _, run_text, bold, italic, color, riscado = seg
        for indice, trecho in enumerate(run_text.split(MARCA_QUEBRA_ODT)):
            if indice:
                p.add_run().add_break()
            if trecho:
                add_text_run(p, trecho, bold, italic, color, is_heading, level, riscado or riscado_inteiro)


def aplicar_classe_sei(doc: DocumentoWord, p, classe: str) -> None:
    """Grava a classe SEI como nome do estilo do parágrafo no DOCX intermediário.

    O DOCX não tem onde guardar uma classe do SEI, e o docx_converter decide a classe
    olhando o parágrafo. O nome do estilo é o único atributo que atravessa inteiro: o
    docx_converter reconhece o nome e usa a classe sem heurística.

    Nome que não é classe do SEI é ignorado: "Strong" é estilo de caractere do Word e
    derrubaria a conversão, e "Heading 1" ou "Quote" mudariam a classe por outro caminho.
    """
    classe_valida = classe_sei_pelo_nome(classe)
    if not classe_valida:
        return
    classe = classe_valida
    try:
        estilo = doc.styles[classe]
    except KeyError:
        estilo = doc.styles.add_style(classe, WD_STYLE_TYPE.PARAGRAPH)
        estilo.base_style = doc.styles["Normal"]
    p.style = estilo


def adicionar_texto_literal(p, texto: str) -> None:
    """Escreve o texto sem ler marcação, com as quebras de linha preservadas.

    Serve ao bloco de código: asterisco, colchete e travessão ali são conteúdo.
    """
    for idx, linha in enumerate(texto.split("\n")):
        if idx:
            p.add_run().add_break()
        if linha:
            p.add_run(linha).font.name = "Calibri"


def comeca_com_item(texto: str) -> bool:
    """Indica se o texto (Markdown) começa com número de item, ignorando as marcas de ênfase."""
    return bool(profundidade_item(desescapar_markdown(texto).lstrip("*~ ")))


def riscar_inteiro(texto: str) -> str:
    """Risca o parágrafo inteiro, sem "~~" dentro de "~~".

    Com o parágrafo todo riscado, a marca de cada trecho é redundante, e "~~a ~~b~~ c~~"
    seria lido como três trechos, com o do meio sem riscado.
    """
    return f"~~{texto.replace('~~', '')}~~"


# Alinhamento lido de ODT que vira classe, a última regra, como no docx_converter. À
# esquerda fica de fora: no ODT ele costuma vir herdado do estilo padrão, e não pedido no
# parágrafo, e tiraria o justificado de quase todo o texto
CLASSE_POR_ALINHAMENTO = {"center": "Texto_Centralizado", "right": "Texto_Alinhado_Direita"}


def linha_com_classe(
    texto: str, formato: Formatacao, explicita: str | None = None, citacao: CitacaoPorRecuo | None = None
) -> str:
    """Linha de Markdown de um parágrafo lido de ODT, com a classe SEI quando houver.

    A classe vai como <p class="...">, que montar_docx_de_markdown grava no DOCX
    intermediário e o docx_converter respeita. O texto é escapado para HTML porque o
    <p> é desmontado e o escape desfeito em seguida. Quando `formato.riscado` é
    verdadeiro, o texto já chega inteiro entre "~~" (riscar_inteiro).

    Um parágrafo que começa com número de item ("1. ASSUNTO") fica sem classe pela
    formatação: quem decide é a detecção de itens do docx_converter, como no Word. Sem
    isso, o título em caixa alta viraria Texto_Justificado_Maiusculas e sairia da
    numeração do SEI. Se estiver riscado, segue item, com o texto riscado. Pelo mesmo
    motivo, a citação pelo recuo (`citacao`) só é considerada depois dessas regras, na
    ordem do docx_converter.
    """
    classe = explicita
    if not classe:
        if comeca_com_item(texto):
            return texto
        visivel = desescapar_markdown(texto)
        formato = replace(formato, maiusculas=formato.maiusculas or texto_em_maiusculas(visivel))
        classe = classe_por_formatacao(formato)
        if not classe and citacao is not None:
            classe = citacao.classe(formato, visivel.replace("~~", ""))
        if not classe:
            classe = CLASSE_POR_ALINHAMENTO.get(formato.alinhamento or "")
    if not classe:
        return texto
    if classe == "Tachado":
        # A classe já risca o parágrafo; o <s> em cada trecho seria redundante
        texto = texto.replace("~~", "")
    return f'<p class="{classe}">{htmlmod.escape(texto, quote=False)}</p>'


def linha_de_assinatura(texto: str, formato: Formatacao) -> str:
    """Linha do bloco de assinatura lido de ODT, com a classe da assinatura.

    A classe não aplica o negrito e a caixa alta que Texto_Centralizado_Maiusculas_Negrito
    aplicava: o texto passa a levá-los, como no docx_converter. A caixa alta do estilo não
    vale para texto com link, cujo endereço mudaria.
    """
    if formato.maiusculas and "](" not in texto:
        texto = texto.upper()
    if formato.negrito and "**" not in texto:
        texto = f"**{texto}**"
    return f'<p class="{CLASSE_ASSINATURA}">{htmlmod.escape(texto, quote=False)}</p>'


def texto_inteiro_riscado(texto: str) -> bool:
    """Indica se todo o texto está entre "~~", em um ou mais trechos."""
    return "~~" in texto and not RE_RISCADO.sub("", texto).strip()


def linha_inteira_riscada(linha: str) -> bool:
    """Indica se a linha é um parágrafo revogado (no SEI, Tachado).

    Todo o texto entre "~~" e sem número de item: um item riscado continua item, para
    não tirar os seguintes da numeração do SEI.
    """
    return texto_inteiro_riscado(linha) and not comeca_com_item(linha)


def neutralizar_marcacao_estendida(linha: str) -> str:
    """Escapa ">", "```" e "~~" para que fiquem como texto.

    Serve ao .txt, que passa pelo mesmo caminho do Markdown mas não é Markdown: em texto
    comum, "> 18 anos" ou uma resposta de e-mail citada não são citação, e "~~" não risca.
    """
    linha = linha.replace("~", "\\~")
    corpo = linha.lstrip()
    if corpo.startswith((">", "```")):
        linha = linha[: len(linha) - len(corpo)] + "\\" + corpo
    return linha


def parse_table_markdown(lines: list[str], start_idx: int) -> tuple[list[list[str]], int]:
    table_data = []
    idx = start_idx

    while idx < len(lines) and lines[idx].strip().startswith("|"):
        line = lines[idx].strip()
        if re.match(r"^\|[\s:\-|\+]+\|$", line):
            idx += 1
            continue

        cells = [c.strip() for c in line.split("|")[1:-1]]
        if cells:
            table_data.append(cells)
        idx += 1

    return table_data, idx


def estilo_de_marcador(doc: DocumentoWord, nivel: int) -> str:
    """Nome do estilo de marcador do nível, criado quando o modelo não o tem.

    O nível atravessa o DOCX intermediário pelo nome ("List Bullet 2"), como no Word,
    e não pelo recuo, que num documento comum não indica subitem.
    """
    nome = f"List Bullet {nivel + 1}" if nivel else "List Bullet"
    if nome not in doc.styles:
        doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["List Bullet"]
    return nome


def fim_do_bloco_de_codigo(linhas: list[str], inicio: int) -> int | None:
    """Índice da cerca que fecha o bloco de código aberto em `inicio`, ou None.

    Fecha com uma linha só de crases, pelo menos tantas quanto as da abertura. Cerca sem
    fechamento é texto: tratá-la como bloco engoliria o resto do documento. É a regra
    única de bloco de código, usada na montagem e nos avisos.
    """
    cerca = RE_CERCA_CODIGO.match(linhas[inicio].strip())
    if not cerca:
        return None
    fechamento = re.compile(rf"^`{{{len(cerca.group(1))},}}$")
    for fim in range(inicio + 1, len(linhas)):
        if fechamento.match(linhas[fim].strip()):
            return fim
    return None


def montar_docx_de_markdown(conteudo: str, estendido: bool = True, extraido: bool = False) -> DocumentoWord:
    """Monta em memória o documento Word institucional (A4, Calibri) a partir do Markdown.

    `estendido` falso (arquivo .txt) deixa como texto as marcações que em texto comum têm
    outro sentido: citação (">"), bloco de código ("```") e riscado ("~~").

    `extraido` verdadeiro indica Markdown gerado pelos leitores de ODT e PDF. Ali o
    "#" de nível 1 é seção do documento, e não o título dele: fica centralizado, como
    sempre ficou, sem a caixa alta de Texto_Centralizado_Maiusculas_Negrito.
    """
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

    # Normaliza o fim de linha antes de fatiar: um documento salvo no Windows (CRLF)
    # deixaria um \r ao fim de cada linha, que vira quebra extra dentro de um bloco
    lines = conteudo.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if not estendido:
        lines = [neutralizar_marcacao_estendida(linha) for linha in lines]
    # A linha de tis só é separador onde "~" é marcação; no .txt, ela é texto
    separador = r"[-–—*_~]{3,}" if estendido else r"[-–—*_]{3,}"
    i = 0

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        fim = fim_do_bloco_de_codigo(lines, i)
        if fim is not None:
            # O conteúdo vai cru, sem strip: a indentação é parte do código
            codigo = "\n".join(linha.rstrip() for linha in lines[i + 1 : fim]).strip("\n")
            if codigo:
                p = doc.add_paragraph()
                adicionar_texto_literal(p, codigo)
                aplicar_classe_sei(doc, p, "Texto_Mono_Espaçado")
            i = fim + 1
            continue

        if re.fullmatch(separador, desescapar_markdown(line)):
            i += 1
            continue

        if line.startswith("<red>"):
            paragrafos, next_idx = iter_red_block_paragraphs(lines, i)
            for texto_red in paragrafos:
                p = doc.add_paragraph()
                format_paragraph(p, f"<red>{texto_red}</red>", is_heading=False)
            i = next_idx
            continue

        html_p = html_paragraph_to_markdown(line)
        if html_p:
            cls_html, text_html = html_p
            p = doc.add_paragraph()
            format_paragraph(
                p,
                text_html,
                is_heading=False,
                is_center="Texto_Centralizado" in cls_html,
                is_left="Texto_Alinhado_Esquerda" in cls_html,
            )
            aplicar_classe_sei(doc, p, cls_html)
            i += 1
            continue

        citacao = RE_CITACAO_MD.match(line)
        if citacao:
            if citacao.group(1).strip():
                p = doc.add_paragraph()
                format_paragraph(p, citacao.group(1).strip(), is_heading=False)
                aplicar_classe_sei(doc, p, "Citação")
            i += 1
            continue

        if linha_inteira_riscada(line):
            # A classe já risca o parágrafo; o <s> em cada trecho seria redundante
            p = doc.add_paragraph()
            format_paragraph(p, line.replace("~~", ""), is_heading=False)
            aplicar_classe_sei(doc, p, "Tachado")
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            # As marcas de ênfase no início ("~~1. X~~" de um título riscado) não escondem o número
            tem_numero = bool(RE_INICIO_NUMERADO.match(desescapar_markdown(text).lstrip("*~ ")))
            p = doc.add_paragraph()
            format_paragraph(p, text, is_heading=True, level=min(level, 4), is_center=not tem_numero)
            if level == 1 and not tem_numero and not extraido:
                # Título do documento, sem número de seção: no SEI, caixa alta em negrito
                aplicar_classe_sei(doc, p, "Texto_Centralizado_Maiusculas_Negrito")
            i += 1
            continue

        if line[:2] in ("- ", "* "):
            text = line[2:].strip()
            tarefa = re.match(r"^\[([ xX])\]\s*(.*)$", text)
            if tarefa:
                text = f"{'Concluído' if tarefa.group(1).lower() == 'x' else 'Pendente'}: {tarefa.group(2)}"
            recuo = len(lines[i]) - len(lines[i].lstrip(" \t"))
            p = doc.add_paragraph(style=estilo_de_marcador(doc, min((recuo + 1) // 2, 3)))
            format_paragraph(p, text, is_heading=False)
            i += 1
            continue

        if line.startswith("|"):
            table_data, next_idx = parse_table_markdown(lines, i)
            if table_data:
                cols_cnt = max(sum(colunas_mescladas(celula)[0] for celula in row) for row in table_data)
                table = doc.add_table(rows=len(table_data), cols=cols_cnt)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                aplicar_bordas_tabela(table)

                for r_idx, row_cells in enumerate(table_data):
                    is_header = r_idx == 0
                    coluna = 0
                    for celula in row_cells:
                        largura, cell_text = colunas_mescladas(celula)
                        c_idx = coluna
                        cell = table.cell(r_idx, c_idx)
                        if largura > 1:
                            cell = cell.merge(table.cell(r_idx, c_idx + largura - 1))
                        definir_margens_celula(cell, top=120, bottom=120, left=150, right=150)
                        if is_header:
                            definir_fundo_celula(cell, "E6E6E6")
                        p = cell.paragraphs[0]
                        format_paragraph(p, cell_text, is_heading=False, is_center=is_header)
                        coluna += largura
                    # Linhas incompletas ainda recebem a largura da tabela.
                    for c_idx in range(coluna, cols_cnt):
                        cell = table.cell(r_idx, c_idx)
                        definir_margens_celula(cell, top=120, bottom=120, left=150, right=150)
                        if is_header:
                            definir_fundo_celula(cell, "E6E6E6")

                i = next_idx
                doc.add_paragraph()
                continue

        p = doc.add_paragraph()
        format_paragraph(p, line, is_heading=False)
        i += 1

    return doc


def markdown_para_docx(md_path: str | Path, docx_path: str | Path) -> Path:
    """Converte arquivo Markdown em documento Word formatado institucionalmente (A4, Calibri)."""
    p_out = Path(docx_path)
    conteudo = decodificar_texto(Path(md_path).read_bytes(), estendido=True)
    doc = montar_docx_de_markdown(conteudo)
    p_out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(p_out))
    return p_out


def converter_texto_md_para_blocos(
    conteudo: str, max_nivel: int = 4, estendido: bool = True, extraido: bool = False
) -> list[str]:
    """Converte texto Markdown em blocos HTML SEI passando o DOCX intermediário em memória.

    O DOCX é etapa interna do pipeline, e não um arquivo que alguém vá abrir, então
    não precisa ir ao disco. Quem entra por texto (o conversor de PDF) chega aqui
    direto, sem escrever nem o .md nem o .docx temporários.
    """
    buffer = io.BytesIO()
    montar_docx_de_markdown(conteudo, estendido, extraido).save(buffer)
    buffer.seek(0)
    return converter_docx_para_blocos(buffer, max_nivel=max_nivel)


def decodificar_texto(conteudo: bytes, estendido: bool = True, avisos: list[str] | None = None) -> str:
    """Reconhece codificações comuns de TXT sem transformar UTF-16 em controles XML.

    Markdown continua com a leitura UTF-8 tolerante anterior. Para texto simples,
    BOM e NULs alternados identificam UTF-16. Um texto que não é UTF-8 válido só é lido
    como Windows-1252 quando quase não tem sequências UTF-8: um byte perdido num arquivo
    UTF-8 viraria mojibake em todos os acentos do documento. Nesse caso, só o byte ruim
    vira U+FFFD, e `avisos` recebe o alerta.
    """
    if conteudo.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return conteudo.decode("utf-16")
    if not estendido and b"\x00" in conteudo:
        pares = len(conteudo) // 2
        zeros_impares = conteudo[1::2].count(0)
        zeros_pares = conteudo[0::2].count(0)
        if pares and zeros_impares > pares // 3 and zeros_pares == 0:
            return conteudo.decode("utf-16-le")
        if pares and zeros_pares > pares // 3 and zeros_impares == 0:
            return conteudo.decode("utf-16-be")
        raise ValueError("O texto contém caracteres de controle. Salve o arquivo como UTF-8 e tente novamente.")
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    if estendido or parece_utf8(conteudo):
        if avisos is not None:
            avisos.append(
                "Alguns caracteres do texto não puderam ser lidos e aparecem como \ufffd. "
                "Confira esses trechos ou salve o arquivo como UTF-8 e converta de novo."
            )
        return conteudo.decode("utf-8-sig", errors="replace")
    try:
        return conteudo.decode("cp1252")
    except UnicodeDecodeError as e:
        raise ValueError("Não foi possível identificar a codificação. Salve o texto como UTF-8.") from e


def parece_utf8(conteudo: bytes) -> bool:
    """Indica se a maior parte dos bytes acima de 127 forma sequências UTF-8 válidas.

    Em Windows-1252, as letras acentuadas do português quase nunca formam essas
    sequências, e em UTF-8 todas formam, menos o byte perdido.
    """
    validos = sum(len(m.group()) for m in RE_SEQUENCIA_UTF8.finditer(conteudo))
    altos = sum(1 for byte in conteudo if byte > 0x7F)
    return validos > altos - validos


def avisar_marcacoes_nao_suportadas(conteudo: str, avisos: list[str]) -> str:
    """Mostra uma marca legível onde uma imagem relativa não pode ser incorporada.

    A seleção de um arquivo Markdown no navegador não concede acesso às imagens da
    mesma pasta. Um link criado a partir de ``![]()`` apontaria para o lugar errado.
    """
    encontrou_imagem = False

    def substituir_imagem(m: re.Match[str]) -> str:
        nonlocal encontrou_imagem
        encontrou_imagem = True
        rotulo = escapar_markdown(m.group(1).strip() or "sem descrição")
        return f"Imagem não incorporada ({rotulo})"

    linhas = conteudo.splitlines(keepends=True)
    saida: list[str] = []
    tem_nota = False
    i = 0
    while i < len(linhas):
        fim = fim_do_bloco_de_codigo(linhas, i)
        if fim is not None:
            # Bloco de código é conteúdo literal: nem imagem nem nota ali dentro
            saida.extend(linhas[i : fim + 1])
            i = fim + 1
            continue
        tem_nota = tem_nota or bool(RE_NOTA_MD.match(linhas[i]))
        saida.append(RE_IMAGEM_MD.sub(substituir_imagem, linhas[i]))
        i += 1
    if encontrou_imagem:
        avisos.append("Imagens indicadas no Markdown não foram incorporadas. Insira-as no SEI após a conversão.")
    if tem_nota:
        avisos.append("Notas de rodapé do Markdown ficaram como texto. Confira as chamadas e o conteúdo das notas.")
    return "".join(saida)


def converter_md_para_blocos(
    md_path: FonteDocumento, max_nivel: int = 4, estendido: bool = True, avisos: list[str] | None = None
) -> list[str]:
    """Converte Markdown (arquivo ou conteúdo em memória) para blocos HTML SEI."""
    conteudo = decodificar_texto(ler_bytes(md_path), estendido=estendido, avisos=avisos)
    if estendido:
        conteudo = avisar_marcacoes_nao_suportadas(conteudo, avisos if avisos is not None else [])
    return converter_texto_md_para_blocos(conteudo, max_nivel=max_nivel, estendido=estendido)


def converter_md_para_html(md_path: FonteDocumento, max_nivel: int = 4) -> str:
    """Converte arquivo Markdown e retorna o HTML do corpo consolidado."""
    blocos = converter_md_para_blocos(md_path, max_nivel=max_nivel)
    return "\n".join(blocos)
